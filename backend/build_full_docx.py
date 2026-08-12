import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def generate_exhaustive_docx():
    doc = Document()

    # Page Margins: 1 inch on all sides
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    COLOR_PRIMARY = RGBColor(15, 81, 50)     # #0F5132 Dark Green
    COLOR_TEXT = RGBColor(20, 24, 27)        # #14181B Dark Charcoal
    COLOR_MUTED = RGBColor(91, 102, 96)      # #5B6660 Gray
    COLOR_HEADER_BG = "0F5132"               # Table header green fill

    def set_cell_bg(cell, hex_color):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    def set_cell_padding(cell, top=100, bottom=100, left=150, right=150):
        tcPr = cell._tc.get_or_add_tcPr()
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        tcPr.append(tcMar)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(18)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = COLOR_TEXT
        return p

    def add_p(text, bold_prefix="", italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = 'Arial'
            rb.font.size = Pt(11)
            rb.font.bold = True
            rb.font.color.rgb = COLOR_TEXT
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(11)
        r.font.italic = italic
        r.font.color.rgb = COLOR_TEXT
        return p

    def add_b(text, bold_prefix=""):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.font.name = 'Arial'
            rb.font.size = Pt(11)
            rb.font.bold = True
            rb.font.color.rgb = COLOR_TEXT
        r = p.add_run(text)
        r.font.name = 'Arial'
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_TEXT
        return p

    def add_code(text):
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = table.cell(0, 0)
        set_cell_bg(cell, "F0F4F2")
        set_cell_padding(cell, top=120, bottom=120, left=180, right=180)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        r.font.name = 'Consolas'
        r.font.size = Pt(9.5)
        r.font.color.rgb = COLOR_TEXT
        doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_tbl(headers, data, widths=None):
        table = doc.add_table(rows=len(data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            set_cell_bg(hdr_cells[i], COLOR_HEADER_BG)
            set_cell_padding(hdr_cells[i], top=120, bottom=120, left=150, right=150)
            p = hdr_cells[i].paragraphs[0]
            for r in p.runs:
                r.font.name = 'Arial'
                r.font.size = Pt(10)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

        for r_idx, row in enumerate(data):
            row_cells = table.rows[r_idx + 1].cells
            bg = "FFFFFF" if r_idx % 2 == 0 else "F8FAF9"
            for c_idx, val in enumerate(row):
                row_cells[c_idx].text = str(val)
                set_cell_bg(row_cells[c_idx], bg)
                set_cell_padding(row_cells[c_idx], top=100, bottom=100, left=150, right=150)
                p = row_cells[c_idx].paragraphs[0]
                for r in p.runs:
                    r.font.name = 'Arial'
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = COLOR_TEXT
        
        if widths:
            for row in table.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Inches(w)
                    
        doc.add_paragraph().paragraph_format.space_after = Pt(6)
        return table

    # =========================================================================
    # 1. TITLE PAGE
    # =========================================================================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(12)
    r_t = p_title.add_run("ParkSafe")
    r_t.font.name = 'Arial'
    r_t.font.size = Pt(32)
    r_t.font.bold = True
    r_t.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    r_s = p_sub.add_run("Smart Vehicle Parking & Online Slot Booking Application\nComplete Technical & Functional Project Documentation")
    r_s.font.name = 'Arial'
    r_s.font.size = Pt(16)
    r_s.font.bold = True
    r_s.font.color.rgb = COLOR_TEXT

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(40)
    p_meta.paragraph_format.space_after = Pt(12)
    r_m = p_meta.add_run("B.Voc (Bachelor of Vocation) Final-Year Project Submission\nDepartment of Software Development / Computer Applications")
    r_m.font.name = 'Arial'
    r_m.font.size = Pt(13)
    r_m.font.bold = True
    r_m.font.color.rgb = COLOR_PRIMARY

    p_team = doc.add_paragraph()
    p_team.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_team.paragraph_format.space_after = Pt(40)
    r_tm = p_team.add_run("PROJECT TEAM:\n• Sayyed Mashhuda\n• Sidat Ibrahim\n• Patel Dhruvil\n\nDate: August 2026")
    r_tm.font.name = 'Arial'
    r_tm.font.size = Pt(12)
    r_tm.font.color.rgb = COLOR_TEXT

    doc.add_page_break()

    # =========================================================================
    # 2. TABLE OF CONTENTS OUTLINE
    # =========================================================================
    add_h1("2. Table of Contents Outline")
    toc_items = [
        "1. Title Page", "2. Table of Contents Outline", "3. Project Overview", "4. Problem Statement",
        "5. Project Objectives", "6. Proposed Solution", "7. User Roles", "8. Technology Stack",
        "9. Project Folder Structure", "10. Application Startup", "11. Authentication", "12. Customer Application",
        "13. Admin Panel", "14. Navigation and Routing", "15. State Management", "16. Repository and Service Layer",
        "17. Database Structure", "18. Database Relationships", "19. Supabase and RLS", "20. Parking Location Management",
        "21. Google Maps", "22. Geolocation", "23. Multi-Floor Parking", "24. Parking Slot Management",
        "25. Booking System", "26. Booking Availability", "27. Booking Validation", "28. Payment",
        "29. QR Code Generation", "30. QR Code Scanning", "31. QR Verification", "32. Check-In",
        "33. Check-Out", "34. Overtime", "35. Notifications", "36. Scheduled Notifications",
        "37. Customer Profile", "38. Refund System", "39. AI Recommendation", "40. Fraud / Overlap Detection",
        "41. Error Handling", "42. Security", "43. API Documentation", "44. Supabase SQL / Functions",
        "45. Package Dependencies", "46. Complete Customer Flow", "47. Complete Admin Flow", "48. Complete Data Flow",
        "49. Testing", "50. Known Limitations", "51. Future Scope", "52. Glossary",
        "53. Viva Questions & Answers", "54. Conclusion"
    ]
    for idx, item in enumerate(toc_items):
        add_b(item)

    doc.add_page_break()

    # =========================================================================
    # 3. PROJECT OVERVIEW
    # =========================================================================
    add_h1("3. Project Overview")
    add_p("ParkSafe is a modern, cross-platform mobile and web application designed to digitize and optimize urban vehicle parking slot reservation and gatekeeping. Developed with Flutter for the client applications and cloud Supabase (PostgreSQL) for backend infrastructure, ParkSafe eliminates the friction of searching for open parking spaces in crowded commercial zones.")
    add_b("Guaranteed Pre-Booking: Drivers can select a destination, date, time window, vehicle type (Car, Bike, EV), and specific floor/slot before traveling.")
    add_b("AI Nearest & Available Recommendation: Combines GPS Haversine distance with real-time slot availability ratio to rank the best parking choices.")
    add_b("Contactless QR Gatekeeping: Generates encrypted digital QR tickets scanned by facility admins for instant gate entry check-in and checkout.")
    add_b("Automated Overtime & Refund Engine: Calculates overtime penalties dynamically upon exit and provides a complete admin workflow for customer refund processing.")

    # =========================================================================
    # 4. PROBLEM STATEMENT
    # =========================================================================
    add_h1("4. Problem Statement")
    add_p("Urban centers face severe transportation bottlenecks due to inefficient manual parking operations:")
    add_b("Time & Fuel Wastage: Motorists spend 15 to 20 minutes cruising for parking, creating 30% of unnecessary city traffic jams and carbon emissions.")
    add_b("Blind Arrival: Drivers arrive at parking garages without knowing if slots are full or whether EV charging points are available.")
    add_b("Paper Ticket Vulnerabilities: Lost paper tokens lead to administrative gate disputes, revenue leakages, and untracked overstaying.")
    add_b("Operational Inefficiency: Garage operators lack real-time visibility into multi-floor occupancy and floor capacity limits.")

    # =========================================================================
    # 5. PROJECT OBJECTIVES
    # =========================================================================
    add_h1("5. Project Objectives")
    add_b("Zero Double Bookings: Implement database-level mathematical overlap query validation to prevent time collisions.")
    add_b("Multi-Floor Visibility: Provide visual, single-page slot layout organized across Floor 1, Floor 2, Floor 3, Basement 1, and Basement 2.")
    add_b("Seamless Mobile Gatekeeping: Deliver camera QR scanning for sub-second gate check-in and check-out verification.")
    add_b("Automated Notifications: Dispatch 15-minute check-in/out reminders and overtime alerts via Firebase FCM and in-app feeds.")

    # =========================================================================
    # 6. PROPOSED SOLUTION
    # =========================================================================
    add_h1("6. Proposed Solution")
    add_p("ParkSafe establishes an end-to-end digital parking ecosystem:")
    add_b("Customer Mobile/Web App: Real-time search, Google Maps direction integration, multi-floor slot selection grid, digital QR pass, My Bookings history, and refund requests.")
    add_b("Admin Operations Suite: Add/edit parking locations, floor capacity management, camera QR scanner, live dashboard statistics, today's bookings list, and customer refund processing.")
    add_b("Python FastAPI AI Microservice: Geospatial distance scoring and background scheduler for automated push notification alerts.")

    # =========================================================================
    # 7. USER ROLES
    # =========================================================================
    add_h1("7. User Roles")
    add_tbl(
        ["Feature / Capability", "Customer Role", "Admin Role"],
        [
            ["Authentication & Login", "Email / Mobile Login", "Admin Credentials"],
            ["Parking Location Search", "View & Search Locations", "Add / Edit / Remove Locations"],
            ["Multi-Floor Slot View", "View Available / Booked Slots", "Configure Floors & Slot Capacities"],
            ["Slot Booking & Payment", "Reserve & Pay for Slots", "View All Customer Bookings"],
            ["QR Ticket Gatekeeping", "View Digital QR Ticket Pass", "Scan QR Code for Check-In / Out"],
            ["Overtime & Late Fees", "View Overtime Fee on Bill", "Automated Checkout Overtime Check"],
            ["Refund Requests", "Submit Refund Request", "Approve / Reject Refund Requests"],
            ["Customer Management", "Edit Personal Profile", "View Registered Customers Directory"]
        ],
        [2.3, 2.1, 2.1]
    )

    # =========================================================================
    # 8. TECHNOLOGY STACK
    # =========================================================================
    add_h1("8. Technology Stack")
    add_tbl(
        ["Technology", "Purpose", "Implementation Status & Where Used"],
        [
            ["Flutter (v3.x)", "UI Framework", "Implemented (Android, Web, iOS Client App)"],
            ["Dart", "Programming Language", "Implemented (Frontend Logic & Provider State)"],
            ["Supabase", "Cloud Backend", "Implemented (PostgreSQL Database, Auth, Realtime)"],
            ["PostgreSQL", "Database Engine", "Implemented (7 Relational Tables & RLS Policies)"],
            ["Firebase FCM", "Push Messaging", "Implemented (Background Notification Service)"],
            ["Provider", "State Management", "Implemented (6 Active ChangeNotifier Classes)"],
            ["GoRouter", "Routing Package", "Available in pubspec.yaml (MaterialApp Routes used)"],
            ["Geolocator", "GPS Location", "Implemented (Customer Location Distance Calc)"],
            ["Google Maps API", "Navigation Link", "Implemented (Direct Turn-by-Turn URLs)"],
            ["QR Flutter", "QR Code Generation", "Implemented (Digital Pass Ticket Rendering)"],
            ["Mobile Scanner", "Camera Scanner", "Implemented (Admin Gate QR Check-In Scanner)"],
            ["FastAPI (Python)", "AI Microservice", "Implemented (Haversine Scoring & API Server)"],
            ["APScheduler", "Background Worker", "Implemented (15-min Reminder Push Scheduler)"],
            ["Figma", "UI/UX Prototyping", "Implemented (Logo Branding & Screen Wireframes)"],
            ["GitHub", "Version Control", "Implemented (Source Repository & Commit Logs)"]
        ],
        [1.8, 1.8, 2.9]
    )

    # =========================================================================
    # 9. COMPLETE PROJECT FOLDER STRUCTURE
    # =========================================================================
    add_h1("9. Complete Project Folder Structure")
    add_code(
"""lib/
├── main.dart                          # App Entry Point, Provider Wiring & Route Map
├── firebase_options.dart              # Firebase Web & Android Credentials
├── core/
│   ├── constants/app_constants.dart   # Supabase URL & Anon Key Constants
│   ├── services/
│   │   ├── notification_service.dart # FCM & Local Push Notification Handler
│   │   └── prediction_service.dart   # FastAPI Recommendation HTTP Client
│   ├── theme/app_theme.dart           # Primary Green (#0F5132) Color Tokens
│   └── utils/
├── data/
│   ├── models/                        # UserProfile, ParkingLocation, ParkingSlot, Booking, etc.
│   ├── repositories/                  # AuthRepository, ParkingRepository, BookingRepository, etc.
│   └── supabase/supabase_client.dart  # Supabase Instance Provider
├── providers/                         # AuthProvider, ParkingProvider, BookingProvider, AdminProvider, etc.
├── screens/
│   ├── admin/                         # AdminShell, Dashboard, AddLocation, Slots, Scanner, Refunds, etc.
│   ├── auth/                          # LoginScreen, RegisterScreen
│   ├── booking/                       # BookingConfirmScreen, PaymentScreen, BookingSuccessScreen
│   ├── home/                          # HomeScreen, Parking Search & Hero Cards
│   ├── location_detail/               # LocationDetailScreen & Facilities Chips
│   ├── location_selection/            # City & Location Selection List
│   ├── my_bookings/                   # MyBookingsScreen, BookingDetailScreen, CustomerRefundsScreen
│   ├── notifications/                 # NotificationsScreen & History Feed
│   ├── onboarding/                    # OnboardingScreen App Tour
│   ├── profile/                       # ProfileScreen & EditProfileScreen
│   ├── recommendation/               # NearestAvailableSlotsScreen (AI List)
│   ├── scanner/                       # ScannerScreen Component
│   ├── slot_selection/                # SlotSelectionScreen (Multi-Floor Grid)
│   └── splash/                        # SplashScreen Session Validation
└── widgets/                           # VehicleNumberPlateField, FeedbackDialog, SlotCard

backend/
├── main.py                            # FastAPI Server & /recommend Endpoint
├── scheduler.py                       # APScheduler Background Notification Worker
├── seed_100_locations.py              # Database Seeder & SQL Generator Script
└── update_100_locations.sql           # Self-Contained SQL Migration File"""
    )

    # =========================================================================
    # 10. APPLICATION STARTUP
    # =========================================================================
    add_h1("10. Application Startup Sequence")
    add_p("The startup flow in lib/main.dart initializes core services synchronously before attaching the UI tree:")
    add_b("1. WidgetsFlutterBinding.ensureInitialized(): Connects Flutter framework engine.")
    add_b("2. Supabase.initialize(): Connects to cloud Supabase PostgreSQL backend using AppConstants.supabaseUrl.")
    add_b("3. NotificationService().init(): Registers Firebase Messaging service worker and local notification channels.")
    add_b("4. SystemChrome UI Styling: Sets system navigation bar colors and locks app orientation to portrait mode.")
    add_b("5. MultiProvider Creation: Wraps MaterialApp with AuthProvider, ParkingProvider, BookingProvider, AdminProvider, NotificationProvider, and RecommendationProvider.")
    add_b("6. SplashScreen Route Execution: Checks current Supabase Auth session. Redirects logged-in customers to /home, admins to /admin, and unauthenticated users to /login.")

    # =========================================================================
    # 11. AUTHENTICATION
    # =========================================================================
    add_h1("11. Authentication Engine")
    add_p("ParkSafe utilizes Supabase Auth integrated with custom user profile storage:")
    add_b("Registration Flow: User inputs full name, email, mobile number, and password. AuthRepository calls Supabase signUp(), creating an auth user and inserting a row into public.profiles with role = 'customer'.")
    add_b("Login Flow: AuthRepository executes signInWithPassword(). Upon success, it fetches public.profiles to verify whether the user role is 'customer' or 'admin'.")
    add_b("Session Persistence: Tokens are securely persisted locally by the Supabase SDK, maintaining active sessions across app restarts.")

    # =========================================================================
    # 12. CUSTOMER APPLICATION
    # =========================================================================
    add_h1("12. Customer Application Suite")
    add_tbl(
        ["Screen Name", "File Path", "Key Capabilities & User Workflow"],
        [
            ["Splash Screen", "lib/screens/splash/splash_screen.dart", "Session check & route dispatching"],
            ["Onboarding Screen", "lib/screens/onboarding/onboarding_screen.dart", "3-slide feature overview for new users"],
            ["Login Screen", "lib/screens/auth/login_screen.dart", "Email/password login form"],
            ["Register Screen", "lib/screens/auth/register_screen.dart", "New customer account registration"],
            ["Home Screen", "lib/screens/home/home_screen.dart", "Location search bar, AI banner & quick links"],
            ["Location Selection", "lib/screens/location_selection/location_selection_screen.dart", "List of 100 parking places with search filter"],
            ["Location Details", "lib/screens/location_detail/location_detail_screen.dart", "Displays price, amenities & Google Maps button"],
            ["Slot Selection", "lib/screens/slot_selection/slot_selection_screen.dart", "Single-page multi-floor grid (Car, Bike, EV)"],
            ["Booking Confirm", "lib/screens/booking/booking_confirm_screen.dart", "Select date, start/end time & number plate"],
            ["Payment Screen", "lib/screens/booking/payment_screen.dart", "Simulated digital payment gateway & invoice"],
            ["Booking Success", "lib/screens/booking/booking_success_screen.dart", "Confirmation badge & digital QR ticket pass"],
            ["My Bookings", "lib/screens/my_bookings/my_bookings_screen.dart", "Tabs for Active Bookings and Past Bookings"],
            ["Booking Detail / QR", "lib/screens/my_bookings/booking_detail_screen.dart", "Displays ticket pass & full scannable QR code"],
            ["Notifications", "lib/screens/notifications/notifications_screen.dart", "15-min check-in/out reminders & overtime alerts"],
            ["Customer Refunds", "lib/screens/my_bookings/customer_refunds_screen.dart", "Customer booking cancellation request form"],
            ["Edit Profile", "lib/screens/profile/edit_profile_screen.dart", "Updates customer full name and mobile number"],
            ["AI Recommended Slots", "lib/screens/recommendation/nearest_available_slots_screen.dart", "Lists nearest parking locations sorted by AI score"]
        ],
        [1.5, 2.5, 2.5]
    )

    # =========================================================================
    # 13. ADMIN PANEL
    # =========================================================================
    add_h1("13. Admin Panel Operational Suite")
    add_tbl(
        ["Admin Screen", "File Path", "Operational Functions"],
        [
            ["Admin Shell", "lib/screens/admin/admin_shell.dart", "Bottom navigation container for admin panel"],
            ["Admin Dashboard", "lib/screens/admin/admin_dashboard_screen.dart", "Displays live occupancy, total revenue & stats"],
            ["Add/Edit Location", "lib/screens/admin/admin_add_location_screen.dart", "Form to add parking location & Google Maps link"],
            ["Slot Configuration", "lib/screens/admin/admin_slots_screen.dart", "Sets total floors & Car/Bike/EV limits per floor"],
            ["QR Gate Scanner", "lib/screens/admin/admin_scan_screen.dart", "Camera QR scanner for gate check-in & check-out"],
            ["Today's Bookings", "lib/screens/admin/admin_todays_bookings_screen.dart", "Lists active check-in/out bookings for today"],
            ["All Bookings", "lib/screens/admin/admin_all_bookings_screen.dart", "Complete searchable booking transaction log"],
            ["Revenue Analytics", "lib/screens/admin/admin_revenue_screen.dart", "Financial revenue breakdowns and occupancy stats"],
            ["Refund Management", "lib/screens/admin/admin_refunds_screen.dart", "Approve or reject customer refund requests"],
            ["Customers Directory", "lib/screens/admin/admin_customers_screen.dart", "Lists registered customer profiles & contact info"]
        ],
        [1.5, 2.5, 2.5]
    )

    # =========================================================================
    # 14. NAVIGATION AND ROUTING
    # =========================================================================
    add_h1("14. Navigation and Routing Map")
    add_p("Navigation is managed using Flutter named routes registered in lib/main.dart:")
    add_code(
"""/                 -> SplashScreen
/onboarding        -> OnboardingScreen
/login             -> LoginScreen
/register          -> RegisterScreen
/location-select   -> LocationSelectionScreen
/home              -> HomeScreen
/admin             -> AdminShell
/notifications     -> NotificationsScreen
/my-refunds        -> CustomerRefundsScreen
/nearest-slots     -> NearestAvailableSlotsScreen"""
    )

    # =========================================================================
    # 15. STATE MANAGEMENT
    # =========================================================================
    add_h1("15. State Management Architecture")
    add_p("ParkSafe implements the Provider pattern with 6 dedicated ChangeNotifier providers:")
    add_b("AuthProvider: Manages login, registration, current user profile, and session logout.")
    add_b("ParkingProvider: Manages location search, location details, and slot loading.")
    add_b("BookingProvider: Manages date/time selection, vehicle number entry, slot reservation, and My Bookings state.")
    add_b("AdminProvider: Handles admin dashboard stats, facility setup, floor configuration, QR scanning, and refund approvals.")
    add_b("NotificationProvider: Manages in-app notification history and read/unread status.")
    add_b("RecommendationProvider: Interacts with FastAPI backend for AI proximity recommendations.")

    # =========================================================================
    # 16. REPOSITORY AND SERVICE LAYER
    # =========================================================================
    add_h1("16. Repository and Service Layer")
    add_p("The data layer is decoupled into repository classes:")
    add_b("AuthRepository: Interacts with Supabase Auth and public.profiles table.")
    add_b("ParkingRepository: Queries public.parking_locations and public.parking_slots.")
    add_b("BookingRepository: Handles booking creation, overlap validation, check-in, check-out, and overtime updates.")
    add_b("AdminRepository: Executes admin location management, slot generation, and analytics queries.")
    add_b("NotificationRepository: Queries and inserts in-app notification records.")
    add_b("RecommendationRepository: HTTP client executing REST requests to FastAPI /recommend endpoint.")

    # =========================================================================
    # 17. DATABASE STRUCTURE
    # =========================================================================
    add_h1("17. Database Structure & Schema")
    add_p("The database consists of 7 PostgreSQL tables hosted on Supabase:")

    add_h2("Table: profiles")
    add_tbl(
        ["Column Name", "Data Type", "Constraints", "Description"],
        [
            ["id", "UUID", "PRIMARY KEY, REFERENCES auth.users", "User unique identifier"],
            ["full_name", "VARCHAR(255)", "NOT NULL", "Customer or admin full name"],
            ["mobile_number", "VARCHAR(20)", "NULLABLE", "Contact phone number"],
            ["fcm_token", "TEXT", "NULLABLE", "Firebase Cloud Messaging token"],
            ["role", "VARCHAR(20)", "DEFAULT 'customer'", "User role ('customer' or 'admin')"],
            ["created_at", "TIMESTAMPTZ", "DEFAULT NOW()", "Account creation timestamp"]
        ],
        [1.5, 1.2, 1.8, 2.0]
    )

    add_h2("Table: parking_locations")
    add_tbl(
        ["Column Name", "Data Type", "Constraints", "Description"],
        [
            ["id", "UUID", "PRIMARY KEY", "Parking place unique ID"],
            ["name", "VARCHAR(255)", "NOT NULL", "Commercial parking name"],
            ["address", "TEXT", "NOT NULL", "Full address string"],
            ["lat", "FLOAT8", "NOT NULL", "Latitude GPS coordinate"],
            ["lng", "FLOAT8", "NOT NULL", "Longitude GPS coordinate"],
            ["total_slots", "INT", "NOT NULL", "Total slot capacity"],
            ["price_per_hour", "NUMERIC(10,2)", "NOT NULL", "Hourly parking rate (₹)"],
            ["rating", "FLOAT8", "DEFAULT 4.5", "Customer rating score"],
            ["amenities", "TEXT[]", "DEFAULT ARRAY[]", "Facilities list (CCTV, EV, etc.)"],
            ["google_maps_link", "TEXT", "NULLABLE", "Direct Google Maps URL"],
            ["image_url", "TEXT", "NULLABLE", "Facility photo URL"]
        ],
        [1.5, 1.2, 1.8, 2.0]
    )

    add_h2("Table: parking_slots")
    add_tbl(
        ["Column Name", "Data Type", "Constraints", "Description"],
        [
            ["id", "UUID", "PRIMARY KEY", "Parking slot unique ID"],
            ["location_id", "UUID", "REFERENCES parking_locations(id)", "Parent parking facility ID"],
            ["slot_number", "VARCHAR(20)", "NOT NULL", "Slot code (e.g. 'C-01', 'B-02')"],
            ["slot_type", "VARCHAR(20)", "NOT NULL", "Vehicle category ('car', 'bike', 'ev')"],
            ["floor", "VARCHAR(50)", "NOT NULL", "Floor level ('Floor 1', 'Basement 1')"],
            ["is_active", "BOOLEAN", "DEFAULT TRUE", "Slot operational status"]
        ],
        [1.5, 1.2, 1.8, 2.0]
    )

    add_h2("Table: bookings")
    add_tbl(
        ["Column Name", "Data Type", "Constraints", "Description"],
        [
            ["id", "UUID", "PRIMARY KEY", "Booking unique ID"],
            ["user_id", "UUID", "REFERENCES profiles(id)", "Customer account ID"],
            ["slot_id", "UUID", "REFERENCES parking_slots(id)", "Reserved parking slot ID"],
            ["booking_date", "DATE", "NOT NULL", "Reservation date (YYYY-MM-DD)"],
            ["start_time", "VARCHAR(10)", "NOT NULL", "Scheduled start time ('14:00')"],
            ["end_time", "VARCHAR(10)", "NOT NULL", "Scheduled end time ('16:00')"],
            ["vehicle_number", "VARCHAR(20)", "NOT NULL", "Vehicle license plate number"],
            ["vehicle_type", "VARCHAR(20)", "NOT NULL", "Vehicle category ('car', 'bike', 'ev')"],
            ["total_amount", "NUMERIC(10,2)", "NOT NULL", "Calculated reservation fee (₹)"],
            ["overtime_fee", "NUMERIC(10,2)", "DEFAULT 0.00", "Late checkout penalty fee (₹)"],
            ["status", "VARCHAR(20)", "DEFAULT 'pending'", "Status ('pending','checked_in','completed')"],
            ["qr_code_data", "TEXT", "NOT NULL", "Encrypted QR pass payload string"]
        ],
        [1.5, 1.2, 1.8, 2.0]
    )

    # =========================================================================
    # 18. DATABASE RELATIONSHIPS
    # =========================================================================
    add_h1("18. Database Relationships")
    add_p("The relational architecture follows clean entity constraints:")
    add_b("profiles (1) ───< (N) bookings: A customer can create multiple bookings.")
    add_b("parking_locations (1) ───< (N) parking_slots: A location contains multiple slots distributed across floors.")
    add_b("parking_slots (1) ───< (N) bookings: A slot can have multiple non-overlapping bookings.")
    add_b("bookings (1) ───< (1) refunds: A booking can have one associated refund request.")

    # =========================================================================
    # 19. SUPABASE AND RLS
    # =========================================================================
    add_h1("19. Supabase and Row Level Security (RLS)")
    add_p("Row Level Security policies protect private customer data while enabling public location discovery:")
    add_b("parking_locations & parking_slots: Public READ permissions for all authenticated customers.")
    add_b("bookings & profiles: Strict RLS policy enforcing auth.uid() = user_id for customer reads/inserts.")
    add_b("Admin Override: Admins possessing role = 'admin' are granted elevated read/write access across all tables.")

    # =========================================================================
    # 20. PARKING LOCATION MANAGEMENT
    # =========================================================================
    add_h1("20. Parking Location Management")
    add_p("Parking facilities are stored in public.parking_locations with full geographic coordinates (lat, lng), address details, price per hour, total capacity, amenities array, and direct Google Maps link.")

    # =========================================================================
    # 21. GOOGLE MAPS
    # =========================================================================
    add_h1("21. Google Maps Navigation Integration")
    add_p("Every location record includes a google_maps_link URL. Tapping 'Get Directions in Google Maps' executes the url_launcher package, launching the external Google Maps application with exact destination coordinates.")

    # =========================================================================
    # 22. GEOLOCATION
    # =========================================================================
    add_h1("22. Geolocation & Distance Calculation")
    add_p("The geolocator package detects the customer's live GPS location. The coordinates are passed to the Haversine distance formula to calculate exact straight-line distance in kilometers and estimated driving ETA.")

    # =========================================================================
    # 23. MULTI-FLOOR PARKING
    # =========================================================================
    add_h1("23. Multi-Floor Parking Architecture")
    add_p("Multi-floor support is implemented using the floor column in public.parking_slots:")
    add_b("Supported Floors: Floor 1, Floor 2, Floor 3, Basement 1, Basement 2.")
    add_b("Admin Setup: Admin defines custom slot counts per floor for Car, Bike, and EV types.")
    add_b("Customer View: Customer app groups slots into floor sections on a single scrollable grid.")

    # =========================================================================
    # 24. PARKING SLOT MANAGEMENT
    # =========================================================================
    add_h1("24. Parking Slot Management")
    add_p("Slots are uniquely identified by slot_number (e.g. C-01, B-02, E-01), slot_type ('car', 'bike', 'ev'), and floor level. Slots carry an is_active boolean flag allowing admins to disable maintenance slots.")

    # =========================================================================
    # 25. BOOKING SYSTEM
    # =========================================================================
    add_h1("25. Booking System")
    add_p("The booking workflow captures customer ID, slot ID, date, time window, vehicle plate number, total price, and generates a unique UUID booking record stored in public.bookings.")

    # =========================================================================
    # 26. BOOKING AVAILABILITY
    # =========================================================================
    add_h1("26. Booking Availability Engine")
    add_p("Slot availability is queried dynamically by evaluating existing non-cancelled bookings for the selected date and slot ID.")

    # =========================================================================
    # 27. BOOKING VALIDATION
    # =========================================================================
    add_h1("27. Booking Validation & Collision Prevention")
    add_p("Double bookings are prevented mathematically:")
    add_code("(NewStart < ExistingEnd) AND (NewEnd > ExistingStart)")
    add_p("If an overlap is detected, the database blocks reservation.")

    # =========================================================================
    # 28. PAYMENT
    # =========================================================================
    add_h1("28. Payment Gateway Workflow")
    add_p("ParkSafe incorporates a simulated digital payment gateway screen that processes credit card, UPI, or net banking confirmation before writing the booking transaction to Supabase.")

    # =========================================================================
    # 29. QR CODE GENERATION
    # =========================================================================
    add_h1("29. QR Code Generation")
    add_p("The qr_flutter package renders a scannable QR code encoding a JSON payload containing booking_id, user_id, vehicle_number, slot_number, and start_time.")

    # =========================================================================
    # 30. QR CODE SCANNING
    # =========================================================================
    add_h1("30. QR Code Scanning")
    add_p("The admin gate scanner uses the mobile_scanner camera feed to read customer QR codes instantly at entry and exit gates.")

    # =========================================================================
    # 31. QR VERIFICATION
    # =========================================================================
    add_h1("31. QR Pass Verification")
    add_p("Decoded QR data is cross-referenced with public.bookings to confirm ticket validity, scheduled start time, and booking status.")

    # =========================================================================
    # 32. CHECK-IN
    # =========================================================================
    add_h1("32. Gate Check-In Logic")
    add_p("Admin confirms vehicle arrival. Status updates to 'checked_in' and actual check-in timestamp is logged. Early check-in before start time is blocked.")

    # =========================================================================
    # 33. CHECK-OUT
    # =========================================================================
    add_h1("33. Gate Check-Out Logic")
    add_p("Admin scans QR code upon vehicle exit. Status transitions to 'completed' ('Past Booking'), and overtime charges are evaluated.")

    # =========================================================================
    # 34. OVERTIME
    # =========================================================================
    add_h1("34. Automated Overtime Engine")
    add_p("Overtime fee formula:")
    add_code("Extra Hours = ceil((ActualCheckoutTime - ScheduledEndTime) / 60)\nOvertime Fee = Extra Hours * Price Per Hour")

    # =========================================================================
    # 35. NOTIFICATIONS
    # =========================================================================
    add_h1("35. Push & In-App Notifications")
    add_p("Firebase Cloud Messaging (FCM) handles background push notifications, while public.notifications stores in-app notification history.")

    # =========================================================================
    # 36. SCHEDULED NOTIFICATIONS
    # =========================================================================
    add_h1("36. Background Scheduled Notifications")
    add_p("The Python FastAPI scheduler runs every 1 minute, checking for 15-minute check-in reminders, 15-minute checkout reminders, and overtime alerts.")

    # =========================================================================
    # 37. CUSTOMER PROFILE
    # =========================================================================
    add_h1("37. Customer Profile Management")
    add_p("Customers can view and update their full name and mobile number in public.profiles using their unique user ID.")

    # =========================================================================
    # 38. REFUND SYSTEM
    # =========================================================================
    add_h1("38. Customer Refund System")
    add_p("Customers can submit cancellation requests prior to check-in. Admins review requests in the Admin Panel and select Approve or Reject.")

    # =========================================================================
    # 39. AI RECOMMENDATION
    # =========================================================================
    add_h1("39. AI Recommendation Microservice")
    add_p("FastAPI backend ranks parking places using a weighted formula:")
    add_code("Final Score = (0.60 * Distance Score) + (0.40 * Availability Score)")

    # =========================================================================
    # 40. FRAUD / OVERLAP DETECTION
    # =========================================================================
    add_h1("40. Fraud & Overlap Detection")
    add_p("Database constraints and strict validation prevent duplicate bookings for the same vehicle plate during overlapping hours.")

    # =========================================================================
    # 41. ERROR HANDLING
    # =========================================================================
    add_h1("41. Global Error Handling")
    add_p("Try-catch blocks wrap all repository network requests, displaying user-friendly error banners upon network timeout or invalid credentials.")

    # =========================================================================
    # 42. SECURITY
    # =========================================================================
    add_h1("42. Security Architecture")
    add_p("All API keys and secrets are securely maintained in environment variables. Sensitive service role keys are [REDACTED] in documentation.")

    # =========================================================================
    # 43. API DOCUMENTATION
    # =========================================================================
    add_h1("43. FastAPI REST API Documentation")
    add_code(
"""POST /recommend
Request Payload:
{
  "user_lat": 21.1702,
  "user_lng": 72.8311,
  "vehicle_type": "car",
  "radius_km": 10.0,
  "limit": 5
}

Response Payload:
[
  {
    "location_id": "uuid",
    "name": "VR Mall Surat",
    "distance_km": 3.42,
    "eta_minutes": 8,
    "free_slots": 14,
    "final_score": 0.82
  }
]"""
    )

    # =========================================================================
    # 44. SUPABASE SQL / FUNCTIONS
    # =========================================================================
    add_h1("44. Supabase SQL Migration Scripts")
    add_p("Self-contained migration script backend/update_100_locations.sql contains CTE queries seeding 100 locations and 1,000 multi-floor slots.")

    # =========================================================================
    # 45. PACKAGE DEPENDENCIES
    # =========================================================================
    add_h1("45. Package Dependencies (pubspec.yaml)")
    add_tbl(
        ["Package Name", "Version", "Primary Purpose"],
        [
            ["supabase_flutter", "^2.3.4", "Cloud Database, Auth & Realtime client"],
            ["provider", "^6.1.2", "State management & reactive UI updates"],
            ["go_router", "^13.2.0", "Declarative navigation routing"],
            ["geolocator", "^14.0.2", "GPS location tracking"],
            ["firebase_messaging", "^15.0.1", "FCM push notification delivery"],
            ["qr_flutter", "^4.1.0", "QR ticket pass code generator"],
            ["mobile_scanner", "^7.4.0", "Camera QR code gate scanner"],
            ["dio", "^5.4.3+1", "HTTP REST API networking client"]
        ],
        [2.0, 1.5, 3.0]
    )

    # =========================================================================
    # 46. COMPLETE CUSTOMER FLOW
    # =========================================================================
    add_h1("46. Complete Customer Journey Flow")
    add_code("Register / Login ──> Search Location ──> Select Vehicle & Date/Time ──> Choose Floor & Slot ──> Payment ──> QR Pass Ticket ──> Gate Check-In ──> Parking ──> Gate Check-Out")

    # =========================================================================
    # 47. COMPLETE ADMIN FLOW
    # =========================================================================
    add_h1("47. Complete Admin Operational Flow")
    add_code("Admin Login ──> Dashboard Overview ──> Manage Locations & Floors ──> Camera QR Scan Gatekeeper ──> Check-In / Out Verification ──> Overtime Check ──> Process Refunds")

    # =========================================================================
    # 48. COMPLETE DATA FLOW
    # =========================================================================
    add_h1("48. End-to-End Data Flow Architecture")
    add_code("Flutter UI Layer <── Provider State <── Repository Client <── Supabase Client <── PostgreSQL Database")

    # =========================================================================
    # 49. TESTING
    # =========================================================================
    add_h1("49. Automated & Manual Test Cases")
    add_tbl(
        ["Test Case ID", "Test Description", "Input Data", "Expected Result", "Status"],
        [
            ["TC-01", "User Login Success", "Valid Credentials", "Redirect to Home Screen", "PASSED"],
            ["TC-02", "Invalid Login", "Incorrect Password", "Display Error Snack Bar", "PASSED"],
            ["TC-03", "Slot Overlap Check", "Overlapping Time", "Database Blocks Reservation", "PASSED"],
            ["TC-04", "QR Pass Scan", "Valid QR Pass Data", "Successful Check-In", "PASSED"],
            ["TC-05", "Overtime Fee Calc", "15 mins Late Checkout", "Adds 1 Hr Overtime Charge", "PASSED"]
        ],
        [1.0, 2.0, 1.5, 1.5, 0.5]
    )

    # =========================================================================
    # 50. KNOWN LIMITATIONS
    # =========================================================================
    add_h1("50. Known Limitations")
    add_b("AI microservice operates as a prototype Python REST API.")
    add_b("Payment gateway operates in test/simulated mode.")
    add_b("FCM push alerts on web browser require service worker HTTPS configuration.")

    # =========================================================================
    # 51. FUTURE SCOPE
    # =========================================================================
    add_h1("51. Future Scope & Roadmap")
    add_b("ANPR automatic license plate barrier camera integration.")
    add_b("IoT ultrasonic parking slot occupation sensors.")
    add_b("AI peak-hour dynamic surge pricing algorithm.")

    # =========================================================================
    # 52. GLOSSARY
    # =========================================================================
    add_h1("52. Technical Glossary")
    add_b("API: Application Programming Interface allowing app components to communicate.")
    add_b("RLS: Row Level Security policies controlling database data access permissions.")
    add_b("FCM: Firebase Cloud Messaging service for mobile push notification delivery.")
    add_b("Provider: Flutter state management pattern for reactive UI updates.")

    # =========================================================================
    # 53. VIVA QUICK REFERENCE (Q&A)
    # =========================================================================
    add_h1("53. Viva Quick Reference (Q&A)")
    qa_list = [
        ("Q: What is ParkSafe?", "A: ParkSafe is a smart parking slot booking application built with Flutter and Supabase that allows pre-booking multi-floor slots."),
        ("Q: How are double bookings prevented?", "A: By enforcing an overlap query: (NewStart < ExistEnd) AND (NewEnd > ExistStart)."),
        ("Q: How does the QR gatekeeper work?", "A: Encodes booking details into a QR code which Admin scans using camera scanner for Check-In and Check-Out.")
    ]
    for q, a in qa_list:
        add_b(q)
        add_p(a, italic=True)

    # =========================================================================
    # 54. CONCLUSION
    # =========================================================================
    add_h1("54. Conclusion")
    add_p("ParkSafe successfully delivers a complete digital parking solution featuring guaranteed pre-booking, multi-floor slot visibility, contactless QR gatekeeping, automated overtime tracking, and AI proximity recommendations.")

    # Save
    p1 = r"C:\Users\isida\Desktop\ParkSafe_Complete_Project_Documentation.docx"
    p2 = r"C:\Users\isida\Desktop\Parksafe\ParkSafe_Complete_Project_Documentation.docx"
    doc.save(p1)
    doc.save(p2)
    print(f"Exhaustive 54-section documentation generated successfully at:\n - {p1}\n - {p2}")

if __name__ == "__main__":
    generate_exhaustive_docx()
