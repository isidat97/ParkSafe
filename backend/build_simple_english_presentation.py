import os
import re
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_simple_english_presentation():
    pptx_path = r'C:/Users/isida/Desktop/ParkSafe_College_Presentation.pptx'
    logo_path = r'C:/Users/isida/Desktop/Parksafe/assets/images/parksafe_logo_clean.png'
    watermark_path = r'C:/Users/isida/Desktop/Parksafe/assets/images/parksafe_watermark.png'

    prs = Presentation()
    prs.slide_width = Inches(13.333) # 16:9 Widescreen
    prs.slide_height = Inches(7.5)

    # Color Palette
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_PRIMARY_GREEN = RGBColor(15, 81, 50) # #0F5132
    COLOR_DARK_TEXT = RGBColor(20, 24, 27) # #14181B
    COLOR_MUTED_TEXT = RGBColor(91, 102, 96) # #5B6660
    COLOR_CARD_BG = RGBColor(248, 250, 249) # Clean white-tint card
    COLOR_BORDER = RGBColor(226, 232, 240) # Light border

    FONT_HEADING = 'Walbaum Display'
    FONT_BODY = 'Aptos Light'

    def set_white_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_WHITE

    def add_watermark(slide):
        if os.path.exists(watermark_path):
            slide.shapes.add_picture(
                watermark_path,
                left=Inches(9.6),
                top=Inches(5.1),
                width=Inches(3.3)
            )

    def add_slide_header(slide, title_text, category_text="PARKSAFE PROJECT PRESENTATION"):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.name = FONT_BODY
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_MUTED_TEXT

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(32) # Main title size
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_GREEN
        p1.space_before = Pt(3)

    def clean_heading_text(title):
        cleaned = re.sub(r'^[^\w\s]+', '', title).strip()
        return cleaned

    def add_card(slide, left, top, width, height, title, points, font_body_size=20, font_title_size=25):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_CARD_BG
        shape.line.color.rgb = COLOR_BORDER
        shape.line.width = Pt(1.5)

        tbox = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        clean_title = clean_heading_text(title)
        p0 = tf.paragraphs[0]
        p0.text = clean_title
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(font_title_size) # 25-28pt
        p0.font.bold = True
        p0.font.color.rgb = COLOR_PRIMARY_GREEN
        p0.space_after = Pt(8)

        for idx, pt in enumerate(points):
            p = tf.add_paragraph()
            p.text = f"• {pt}"
            p.font.name = FONT_BODY
            p.font.size = Pt(font_body_size) # 20-22pt perfectly calculated to stay inside card
            p.font.color.rgb = COLOR_DARK_TEXT
            p.space_after = Pt(6)
            p.line_spacing = 1.12

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide1)

    if os.path.exists(logo_path):
        slide1.shapes.add_picture(logo_path, Inches(0.8), Inches(0.6), width=Inches(4.2))

    tbox1 = slide1.shapes.add_textbox(Inches(0.8), Inches(2.1), Inches(11.7), Inches(2.4))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

    p_title = tf1.paragraphs[0]
    p_title.text = "ParkSafe"
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(48)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_PRIMARY_GREEN

    p_sub = tf1.add_paragraph()
    p_sub.text = "Smart Vehicle Parking & Online Slot Booking Application"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(24)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_DARK_TEXT
    p_sub.space_before = Pt(6)

    p_tag = tf1.add_paragraph()
    p_tag.text = "FIND. BOOK. PARK SAFE."
    p_tag.font.name = FONT_BODY
    p_tag.font.size = Pt(16)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_PRIMARY_GREEN
    p_tag.space_before = Pt(6)

    team_box = slide1.shapes.add_textbox(Inches(0.8), Inches(4.7), Inches(5.5), Inches(2.2))
    tf_team = team_box.text_frame
    tf_team.word_wrap = True

    pt0 = tf_team.paragraphs[0]
    pt0.text = "TEAM MEMBERS:"
    pt0.font.name = FONT_BODY
    pt0.font.size = Pt(16)
    pt0.font.bold = True
    pt0.font.color.rgb = COLOR_PRIMARY_GREEN
    pt0.space_after = Pt(6)

    for name in ["Sayyed Mashhuda", "Sidat Ibrahim", "Patel Dhruvil"]:
        pt = tf_team.add_paragraph()
        pt.text = f"• {name}"
        pt.font.name = FONT_BODY
        pt.font.size = Pt(18)
        pt.font.color.rgb = COLOR_DARK_TEXT
        pt.space_after = Pt(4)

    dept_box = slide1.shapes.add_textbox(Inches(6.8), Inches(4.7), Inches(5.5), Inches(2.2))
    tf_dept = dept_box.text_frame
    tf_dept.word_wrap = True

    pd0 = tf_dept.paragraphs[0]
    pd0.text = "DEPARTMENT:"
    pd0.font.name = FONT_BODY
    pd0.font.size = Pt(16)
    pd0.font.bold = True
    pd0.font.color.rgb = COLOR_PRIMARY_GREEN
    pd0.space_after = Pt(6)

    pd1 = tf_dept.add_paragraph()
    pd1.text = "B.Voc (Bachelor of Vocation)"
    pd1.font.name = FONT_BODY
    pd1.font.size = Pt(20)
    pd1.font.bold = True
    pd1.font.color.rgb = COLOR_DARK_TEXT
    pd1.space_after = Pt(4)

    pd2 = tf_dept.add_paragraph()
    pd2.text = "Software Development / Computer Applications"
    pd2.font.name = FONT_BODY
    pd2.font.size = Pt(16)
    pd2.font.color.rgb = COLOR_MUTED_TEXT

    # =========================================================================
    # SLIDE 2: INTRODUCTION & PROJECT BACKGROUND
    # =========================================================================
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide2)
    add_watermark(slide2)
    add_slide_header(slide2, "1. Introduction & Project Background")

    add_card(slide2, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.3),
             "The Urban Parking Problem",
             ["City roads are getting crowded with more cars and bikes every day.",
              "Drivers waste a lot of time searching for empty parking spaces.",
              "Traffic jams increase because vehicles keep looking for open slots."],
             font_body_size=20, font_title_size=25)

    add_card(slide2, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.3),
             "The Shift to Digital Booking",
             ["Traditional paper tickets are slow and easy to lose.",
              "Pre-booking parking slots online removes stress before a journey.",
              "Customers can check real-time slot availability on their phones."],
             font_body_size=20, font_title_size=25)

    add_card(slide2, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.3),
             "The ParkSafe Vision",
             ["A complete mobile app built with Flutter and Supabase cloud.",
              "AI helps find parking locations likely to have free slots.",
              "Fast QR code scanning for quick gate check-in and check-out."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 3: PROBLEM STATEMENT
    # =========================================================================
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide3)
    add_watermark(slide3)
    add_slide_header(slide3, "2. Problem Statement: Urban Parking Crises")

    add_card(slide3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.55),
             "Time & Fuel Wastage",
             ["Drivers spend 15 to 20 minutes just looking for a parking spot.",
              "This searching creates up to 30% of traffic jams in busy areas."],
             font_body_size=20, font_title_size=25)

    add_card(slide3, Inches(6.8), Inches(1.6), Inches(5.7), Inches(2.55),
             "Lack of Real-Time Info",
             ["People drive to parking places without knowing if slots are full.",
              "No clear information for Cars, Bikes, or EV Charging stations."],
             font_body_size=20, font_title_size=25)

    add_card(slide3, Inches(0.8), Inches(4.35), Inches(5.6), Inches(2.55),
             "Paper Ticket Problems",
             ["Lost paper tickets cause long arguments and gate delays.",
              "Extra time charges are hard to track without a digital timer."],
             font_body_size=20, font_title_size=25)

    add_card(slide3, Inches(6.8), Inches(4.35), Inches(5.7), Inches(2.55),
             "Security & Verification",
             ["Hard to verify ticket ownership without digital records.",
              "Drivers get no reminders when their parking time is ending."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 4: EXISTING PROBLEMS VS PARKSAFE SOLUTION
    # =========================================================================
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide4)
    add_watermark(slide4)
    add_slide_header(slide4, "3. Existing Problems vs. ParkSafe Solution")

    add_card(slide4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3),
             "Traditional Manual Parking",
             ["Blind Arrival: Drivers arrive and hope to find an open spot.",
              "Slow Paper Tickets: Paper slips cause long lines at entrance gates.",
              "No Floor Visibility: Managers cannot see live slot counts per floor.",
              "Untracked Overtime: Vehicles staying late are rarely charged correctly."],
             font_body_size=20, font_title_size=25)

    add_card(slide4, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3),
             "ParkSafe Smart Solution",
             ["Guaranteed Pre-Booking: Reserve parking slots before leaving home.",
              "Digital QR Pass: Fast QR code scanning on smartphone.",
              "Multi-Floor Grid: Clear view of open slots floor by floor.",
              "Automatic Overtime: Late fees calculated automatically upon exit."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 5: PROPOSED SOLUTION - THE PARKSAFE ECOSYSTEM
    # =========================================================================
    slide5 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide5)
    add_watermark(slide5)
    add_slide_header(slide5, "4. Proposed Solution: The ParkSafe Ecosystem")

    add_card(slide5, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.3),
             "Easy Pre-Booking",
             ["Book parking slots in advance for specific dates and times.",
              "Filter slots for Cars, Bikes, or EV Charging stations.",
              "View available floor slots on a single simple screen."],
             font_body_size=20, font_title_size=25)

    add_card(slide5, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.3),
             "AI Parking Recommender",
             ["Uses GPS to measure exact distance to nearby locations.",
              "Checks real-time slot availability to suggest best options.",
              "Fast Python FastAPI backend calculates top recommendations."],
             font_body_size=20, font_title_size=25)

    add_card(slide5, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.3),
             "QR Gatekeeper & Alerts",
             ["Instant QR code scanning for entry and check-out.",
              "Automatic overtime calculation if a driver leaves late.",
              "Phone push notifications for check-in and check-out reminders."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 6: PROJECT OBJECTIVES & SCOPE
    # =========================================================================
    slide6 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide6)
    add_watermark(slide6)
    add_slide_header(slide6, "5. Key Objectives & Project Scope")

    add_card(slide6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3),
             "Main Technical Goals",
             ["No Double Booking: Database blocks two users from booking the same slot at the same time.",
              "Multi-Floor Support: Manage parking slots across multiple floor levels easily.",
              "Fast Cloud Sync: Instant update between Customer App and Admin Panel.",
              "Multi-Platform: Runs smoothly on Android, iOS, and Web browsers."],
             font_body_size=20, font_title_size=25)

    add_card(slide6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3),
             "Project Scope & Boundaries",
             ["Customer App: Location search, AI recommendation, slot booking, QR ticket.",
              "Admin Panel: Add locations, set floor slots, QR scanner, view live earnings.",
              "Python Backend: AI recommendation backend and automated notification worker."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 7: KEY FEATURES & CAPABILITIES OVERVIEW
    # =========================================================================
    slide7 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide7)
    add_watermark(slide7)
    add_slide_header(slide7, "6. Key Features & Capabilities Overview")

    features = [
        ("Multi-Floor Setup", "Set custom floor layouts and slot limits for Car, Bike, and EV."),
        ("AI Recommender", "Ranks locations based on short distance and available free slots."),
        ("Overlap Protection", "Database checks time slots so two people cannot book the same slot."),
        ("QR Pass Scanner", "Quick camera scanning at gates with early check-in protection."),
        ("Push Notifications", "Automated reminders 15 minutes before check-in and check-out."),
        ("Overtime & Refunds", "Automatic late fee calculation and easy customer refund requests.")
    ]

    for idx, (ftitle, fdesc) in enumerate(features):
        r = idx // 3
        c = idx % 3
        left = Inches(0.8 + c * 3.9)
        top = Inches(1.6 + r * 2.7)
        add_card(slide7, left, top, Inches(3.6), Inches(2.55), ftitle, [fdesc], font_body_size=19, font_title_size=24)

    # =========================================================================
    # SLIDE 8: CUSTOMER APPLICATION FEATURES
    # =========================================================================
    slide8 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide8)
    add_watermark(slide8)
    add_slide_header(slide8, "7. Customer Application Features")

    add_card(slide8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3),
             "Search & Location Discovery",
             ["Simple user registration and secure login.",
              "Search nearby parking places with Google Maps directions.",
              "View facilities like 24/7 CCTV, EV Charging, and Valet Parking.",
              "Get AI suggestions for nearest available parking spots."],
             font_body_size=20, font_title_size=25)

    add_card(slide8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3),
             "Slot Reservation & Digital Pass",
             ["Select booking date, start time, and end time.",
              "Filter slots for Cars, Bikes, or EV Charging stations.",
              "Choose floor and slot number on an interactive visual grid.",
              "Get an instant QR ticket pass saved under My Bookings."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 9: ADMIN PANEL OPERATIONAL SUITE FEATURES
    # =========================================================================
    slide9 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide9)
    add_watermark(slide9)
    add_slide_header(slide9, "8. Admin Panel Operational Suite Features")

    add_card(slide9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3),
             "Facility & Slot Configuration",
             ["Add or update parking locations with Google Maps links.",
              "Set total floors (Floor 1, Floor 2, Basement) and slot counts.",
              "Set separate slot limits for Cars, Bikes, and EV vehicles.",
              "View live dashboard with total earnings and today's bookings."],
             font_body_size=20, font_title_size=25)

    add_card(slide9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3),
             "Gate Operation & Refunds",
             ["Fast camera QR scanner to check vehicles in and out.",
              "Search by vehicle number plate if QR scanning fails.",
              "Blocks check-in if customer arrives before start time.",
              "Review and approve or reject customer refund requests."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 10: PARKING LOCATION & MULTI-FLOOR MANAGEMENT
    # =========================================================================
    slide10 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide10)
    add_watermark(slide10)
    add_slide_header(slide10, "9. Multi-Floor Parking & Slot Management")

    add_card(slide10, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3),
             "Admin Multi-Floor Setup",
             ["Flexible Floor Setup: Each location can have different number of floors.",
              "Independent Floor Limits: E.g., Floor 1 (20 slots), Floor 2 (25 slots).",
              "Vehicle Allocation: Specify Car, Bike, and EV slots for each floor.",
              "Auto Slot Generator: System creates slot codes like C-01 or E-01 automatically."],
             font_body_size=20, font_title_size=25)

    add_card(slide10, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3),
             "Customer Floor Layout",
             ["Single-Page Screen: All floors displayed together on one screen.",
              "Clear Floor Headers: Sections labeled Floor 1, Floor 2, Basement 1.",
              "Live Available Counters: Displays free slots available on each floor.",
              "Clear Ticket Details: Ticket shows exact location e.g., Floor 2 -> Slot C-24."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 11: SLOT BOOKING & REAL-TIME AVAILABILITY SYSTEM
    # =========================================================================
    slide11 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide11)
    add_watermark(slide11)
    add_slide_header(slide11, "10. Slot Booking & Real-Time Availability")

    add_card(slide11, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3),
             "Real-Time Database Sync",
             ["System checks active bookings before confirming a slot reservation.",
              "Visual Color Badges: Available slots show Green, Booked slots show Red.",
              "Instant Update: Cancelled bookings immediately open slots for others.",
              "Cloud Synchronization: Changes reflect instantly on all devices."],
             font_body_size=20, font_title_size=25)

    add_card(slide11, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3),
             "Same-Day Multiple Bookings",
             ["Multiple users can book the same slot at different times on the same day.",
              "Example: Slot C-01 booked from 1:00 PM to 2:00 PM and 3:00 PM to 4:00 PM.",
              "Prevents time collisions while maximizing slot usage throughout the day."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 12: BOOKING VALIDATION & OVERLAPPING BOOKING PREVENTION
    # =========================================================================
    slide12 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide12)
    add_watermark(slide12)
    add_slide_header(slide12, "11. Booking Validation & Overlap Prevention")

    add_card(slide12, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3),
             "Database Time Overlap Rule",
             ["Checks if new booking time overlaps with existing booking time.",
              "Overlap Condition: (NewStart < ExistEnd) AND (NewEnd > ExistStart).",
              "If times overlap, the database blocks the booking attempt immediately.",
              "Ensures zero double bookings for all vehicle types."],
             font_body_size=20, font_title_size=25)

    add_card(slide12, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3),
             "Clear Example",
             ["Existing Booking: 12 August (1:00 PM to 2:00 PM).",
              "ALLOWED: 12 August (3:00 PM to 4:00 PM) -> No time overlap.",
              "BLOCKED: 12 August (1:30 PM to 2:30 PM) -> Overlaps by 30 minutes.",
              "Rule: Prevents double booking for the same slot at overlapping times."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 13: QR-BASED CHECK-IN & CHECK-OUT VERIFICATION FLOW
    # =========================================================================
    slide13 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide13)
    add_watermark(slide13)
    add_slide_header(slide13, "12. QR-Based Check-In & Check-Out Flow")

    add_card(slide13, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3),
             "Step-by-Step Gate Workflow",
             ["1. Customer books slot -> QR ticket generated on phone.",
              "2. Customer arrives at gate -> Admin scans QR code with camera.",
              "3. Check-In: Booking status changes to 'Active / Checked In'.",
              "4. Customer leaves -> Admin scans QR again to perform Check-Out.",
              "5. Status changes from 'Active' to 'Past' in My Bookings."],
             font_body_size=20, font_title_size=25)

    add_card(slide13, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3),
             "Early Arrival & Manual Backup",
             ["Early Check-In Block: Prevents gate entry before scheduled start time.",
              "Manual Plate Lookup: Admin can enter vehicle number if camera fails.",
              "Time Recording: Stores exact check-in and check-out times in database."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 14: NOTIFICATIONS, REMINDERS & OVERTIME MANAGEMENT
    # =========================================================================
    slide14 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide14)
    add_watermark(slide14)
    add_slide_header(slide14, "13. Notifications & Overtime Management")

    add_card(slide14, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3),
             "Real-Time Push Notifications",
             ["Check-In Reminder: Sent 15 minutes before parking start time.",
              "Check-Out Reminder: Sent 15 minutes before parking end time.",
              "Overtime Warning: Sent if vehicle stays past the booked end time.",
              "Notification History: Saved in database and visible inside customer app."],
             font_body_size=20, font_title_size=25)

    add_card(slide14, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3),
             "Automated Overtime Calculation",
             ["Compares Scheduled Check-Out Time with Actual Exit Time.",
              "Example: Scheduled 10:00 PM, Actual Exit 10:15 PM (15 mins late).",
              "Round Up Hours: Extra Hours = ceil(overtime_mins / 60) -> 1 hour.",
              "Late Fee: Overtime Fee = Extra Hours * Hourly Rate added to final bill."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 15: REFUND & CUSTOMER ACCOUNT MANAGEMENT
    # =========================================================================
    slide15 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide15)
    add_watermark(slide15)
    add_slide_header(slide15, "14. Refund & Customer Account Management")

    add_card(slide15, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3),
             "Customer Refund Workflow",
             ["Easy Request: Customer can request cancellation before check-in.",
              "Admin Review: Admin sees refund request and reason on dashboard.",
              "Action: Admin clicks Approve or Reject button.",
              "Instant Update: Customer receives push notification with refund status."],
             font_body_size=20, font_title_size=25)

    add_card(slide15, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3),
             "Customer Profile Editing",
             ["Easy Profile Edit: Customer can update Full Name and Mobile Number.",
              "No Avatar Upload: Focused only on basic account information.",
              "Database Update: Updates existing record in Supabase using User ID.",
              "Data Sync: Updated info shows instantly in Customer App and Admin Panel."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 16: AI NEAREST & AVAILABLE PARKING RECOMMENDATION
    # =========================================================================
    slide16 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide16)
    add_watermark(slide16)
    add_slide_header(slide16, "15. AI Nearest & Available Recommendation")

    add_card(slide16, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.3),
             "Distance Calculation",
             ["Uses GPS coordinates to measure exact distance in kilometers.",
              "Estimates driving ETA time in minutes based on distance."],
             font_body_size=20, font_title_size=25)

    add_card(slide16, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.3),
             "Availability Scoring",
             ["Checks live free slots for requested vehicle type (Car, Bike, EV).",
              "Availability Score = Free Slots / Total Slots.",
              "Hides parking locations that have no empty slots available."],
             font_body_size=20, font_title_size=25)

    add_card(slide16, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.3),
             "Simple Score & Status",
             ["Final Score = 0.60 * Distance Score + 0.40 * Availability Score.",
              "Ranks top recommendations so user gets close & open parking.",
              "Status: Implemented via Python FastAPI microservice (Prototype)."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 17: SYSTEM ARCHITECTURE & TECHNOLOGY STACK
    # =========================================================================
    slide17 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide17)
    add_watermark(slide17)
    add_slide_header(slide17, "16. System Architecture & Tech Stack")

    add_card(slide17, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3),
             "Layered System Architecture",
             ["Frontend: Flutter mobile and web app built using Dart language.",
              "State Management: Provider pattern for smooth screen updates.",
              "Backend Database: Supabase PostgreSQL with secure access policies.",
              "Push Server: Firebase Cloud Messaging (FCM) for notifications.",
              "Python Microservice: FastAPI server for AI scoring & background timers."],
             font_body_size=20, font_title_size=25)

    add_card(slide17, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3),
             "External Services & Tools",
             ["Google Maps API: Opens direct location navigation links.",
              "Geolocator: Detects user's GPS position for distance calculation.",
              "Mobile Scanner: High-speed camera QR code reader package.",
              "Figma & GitHub: Screen design and source code version control."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 18: DATABASE SCHEMA & SUPABASE BACKEND
    # =========================================================================
    slide18 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide18)
    add_watermark(slide18)
    add_slide_header(slide18, "17. Database Schema & Supabase Backend")

    add_card(slide18, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3),
             "Main Database Tables",
             ["Profiles: Stores user ID, full name, mobile number, and FCM token.",
              "Parking Locations: Stores location name, address, GPS coords, and price.",
              "Parking Slots: Stores slot code, vehicle type (Car, Bike, EV), and floor.",
              "Bookings: Stores user ID, slot ID, date, time range, status, and late fee."],
             font_body_size=20, font_title_size=25)

    add_card(slide18, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3),
             "Table Relationships",
             ["Location contains multiple Floors and Slots.",
              "Slot belongs to a specific Parking Location.",
              "Booking connects Customer, Parking Location, and Slot.",
              "Row Level Security (RLS) protects user data from unauthorized access."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 19: COMPLETE BOOKING WORKFLOW / SYSTEM FLOW
    # =========================================================================
    slide19 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide19)
    add_watermark(slide19)
    add_slide_header(slide19, "18. Complete System End-to-End Workflow")

    add_card(slide19, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3),
             "Customer Journey",
             ["1. User logs into ParkSafe app.",
              "2. Searches location and views Google Maps directions.",
              "3. Selects vehicle type (Car, Bike, EV) and date/time range.",
              "4. Chooses floor and slot on visual grid screen.",
              "5. Enters vehicle number, pays, and gets QR ticket pass."],
             font_body_size=20, font_title_size=25)

    add_card(slide19, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3),
             "Admin Journey",
             ["6. Admin scans customer QR code at entry gate.",
              "7. System verifies pass and checks vehicle in.",
              "8. Customer parks vehicle inside garage.",
              "9. Admin scans QR code at exit gate to check out.",
              "10. System checks for late fees and moves booking to Past."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 20: BENEFITS, FUTURE SCOPE & CONCLUSION
    # =========================================================================
    slide20 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide20)
    add_watermark(slide20)
    add_slide_header(slide20, "19. Benefits, Future Scope & Conclusion")

    add_card(slide20, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.3),
             "Main Benefits & Impact",
             ["Saves Time: Eliminates 15-20 minutes of hunting for parking.",
              "Paperless: 100% digital QR code passes on mobile phone.",
              "Easy Management: Automated gate check-ins and overtime tracking.",
              "Eco-Friendly: Reduces unnecessary driving and car emissions."],
             font_body_size=20, font_title_size=25)

    add_card(slide20, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.3),
             "Future Enhancements & Conclusion",
             ["Camera Barrier: Automatic License Plate Recognition at gate entry.",
              "Smart Sensors: Ultrasonic IoT sensors for automatic slot detection.",
              "Dynamic Pricing: Price adjustments during high-demand peak hours.",
              "Summary: Complete digital solution for smart city parking."],
             font_body_size=20, font_title_size=25)

    # =========================================================================
    # SLIDE 21: REFERENCES & TECHNOLOGIES USED
    # =========================================================================
    slide_ref = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide_ref)
    add_watermark(slide_ref)
    add_slide_header(slide_ref, "References & Technologies Used", category_text="PARKSAFE - PROJECT ARCHITECTURE & REFERENCES")

    tech_stack = [
        ("Flutter", "Cross-platform framework for Web, Android & iOS apps"),
        ("Dart", "Primary programming language for mobile application"),
        ("Supabase", "Cloud database, user authentication & realtime backend"),
        ("PostgreSQL", "Relational database engine with secure data policies"),
        ("Firebase FCM", "Push notification service for reminders & alerts"),
        ("Google Maps API", "Map location links and turn-by-turn navigation"),
        ("Geolocator", "GPS tracking to measure distance to parking"),
        ("Python & FastAPI", "AI recommendation microservice REST API"),
        ("APScheduler", "Background worker for scheduled notifications"),
        ("Figma", "UI/UX screen design and logo branding assets"),
        ("Antigravity", "AI-assisted development workspace & code intelligence"),
        ("GitHub", "Version control and code repository management")
    ]

    for idx, (tech_name, tech_desc) in enumerate(tech_stack):
        r = idx // 3
        c = idx % 3
        left = Inches(0.8 + c * 3.9)
        top = Inches(1.6 + r * 1.35)
        
        shape = slide_ref.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.6), Inches(1.2))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_CARD_BG
        shape.line.color.rgb = COLOR_BORDER
        shape.line.width = Pt(1.2)

        tbox = slide_ref.shapes.add_textbox(left + Inches(0.15), top + Inches(0.12), Inches(3.3), Inches(0.95))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = tech_name
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(18)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_PRIMARY_GREEN

        p1 = tf.add_paragraph()
        p1.text = tech_desc
        p1.font.name = FONT_BODY
        p1.font.size = Pt(12)
        p1.font.color.rgb = COLOR_DARK_TEXT
        p1.space_before = Pt(2)

    # =========================================================================
    # SLIDE 22 (FINAL): THANK YOU SLIDE
    # =========================================================================
    slide22 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide22)
    add_watermark(slide22)

    tbox22 = slide22.shapes.add_textbox(Inches(1.5), Inches(1.2), Inches(10.3), Inches(5.2))
    tf22 = tbox22.text_frame
    tf22.word_wrap = True

    p_t0 = tf22.paragraphs[0]
    p_t0.alignment = PP_ALIGN.CENTER
    p_t0.text = "PARKSAFE PROJECT PRESENTATION"
    p_t0.font.name = FONT_BODY
    p_t0.font.size = Pt(14)
    p_t0.font.bold = True
    p_t0.font.color.rgb = COLOR_MUTED_TEXT

    p_t1 = tf22.add_paragraph()
    p_t1.alignment = PP_ALIGN.CENTER
    p_t1.text = "Thank You!"
    p_t1.font.name = FONT_HEADING
    p_t1.font.size = Pt(48)
    p_t1.font.bold = True
    p_t1.font.color.rgb = COLOR_PRIMARY_GREEN
    p_t1.space_before = Pt(10)

    p_t2 = tf22.add_paragraph()
    p_t2.alignment = PP_ALIGN.CENTER
    p_t2.text = "Smart Vehicle Parking & Online Slot Booking Application"
    p_t2.font.name = FONT_BODY
    p_t2.font.size = Pt(22)
    p_t2.font.bold = True
    p_t2.font.color.rgb = COLOR_DARK_TEXT
    p_t2.space_before = Pt(6)

    p_t3 = tf22.add_paragraph()
    p_t3.alignment = PP_ALIGN.CENTER
    p_t3.text = "Questions & Evaluation"
    p_t3.font.name = FONT_HEADING
    p_t3.font.size = Pt(28)
    p_t3.font.bold = True
    p_t3.font.color.rgb = COLOR_PRIMARY_GREEN
    p_t3.space_before = Pt(24)

    p_t4 = tf22.add_paragraph()
    p_t4.alignment = PP_ALIGN.CENTER
    p_t4.text = "Team: Sayyed Mashhuda  |  Sidat Ibrahim  |  Patel Dhruvil"
    p_t4.font.name = FONT_BODY
    p_t4.font.size = Pt(16)
    p_t4.font.color.rgb = COLOR_DARK_TEXT
    p_t4.space_before = Pt(12)

    p_t5 = tf22.add_paragraph()
    p_t5.alignment = PP_ALIGN.CENTER
    p_t5.text = "Department of B.Voc (Bachelor of Vocation) | Software Development"
    p_t5.font.name = FONT_BODY
    p_t5.font.size = Pt(14)
    p_t5.font.color.rgb = COLOR_MUTED_TEXT
    p_t5.space_before = Pt(6)

    prs.save(pptx_path)
    print(f"Simple English presentation saved successfully with {len(prs.slides)} slides to: {pptx_path}")

if __name__ == "__main__":
    build_simple_english_presentation()
