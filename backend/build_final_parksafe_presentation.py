import os
import re
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def generate_presentation():
    input_pptx = r'C:\Users\isida\Desktop\Parksafe\ParkSafe_College_Presentation.pptx'
    output_pptx = r'C:\Users\isida\Desktop\Parksafe\ParkSafe_College_Presentation.pptx'
    img_folder = r'C:\Users\isida\Pictures\Project'

    prs = Presentation(input_pptx)

    # 1. Truncate slides after index 21 (Keep Slides 1..22 untouched!)
    print(f"Initial total slides: {len(prs.slides)}")
    while len(prs.slides) > 22:
        rId = prs.slides._sldIdLst[22].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[22]
    print(f"Slides after keeping protected Slides 1..22: {len(prs.slides)}")

    # Design Constants
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_PRIMARY_GREEN = RGBColor(15, 81, 50)  # #0F5132
    COLOR_DARK_TEXT = RGBColor(20, 24, 27)      # #14181B
    COLOR_MUTED_TEXT = RGBColor(91, 102, 96)    # #5B6660
    COLOR_CARD_BG = RGBColor(248, 250, 249)     # White-tint card
    COLOR_BORDER = RGBColor(226, 232, 240)      # Light border

    FONT_HEADING = 'Walbaum Display'
    FONT_BODY = 'Aptos Light'

    def set_white_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_WHITE

    def add_slide_header(slide, title_text, category_text="PARKSAFE PROJECT DEMONSTRATION"):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.name = FONT_BODY
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_MUTED_TEXT

        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(30)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_GREEN
        p1.space_before = Pt(3)

    def add_screenshot_slide(prs, img_filename, title, description, category="CUSTOMER APP DEMONSTRATION"):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        set_white_background(slide)
        add_slide_header(slide, title, category_text=category)

        img_path = os.path.join(img_folder, img_filename)
        if not os.path.exists(img_path):
            print(f"Warning: image {img_filename} not found!")
            return slide

        # Add background card frame for screenshot
        card_width = Inches(2.45)
        card_height = Inches(5.15)
        card_left = Inches((13.333 - 2.45) / 2) # Center horizontally
        card_top = Inches(1.42)

        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card_left, card_top, card_width, card_height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_CARD_BG
        shape.line.color.rgb = COLOR_BORDER
        shape.line.width = Pt(1.5)

        # Place Screenshot Image inside frame
        img_width = Inches(2.25)
        img_height = Inches(5.0)
        img_left = Inches((13.333 - 2.25) / 2)
        img_top = Inches(1.495)

        slide.shapes.add_picture(img_path, img_left, img_top, width=img_width, height=img_height)

        # Add one-line description text below screenshot
        desc_box = slide.shapes.add_textbox(Inches(1.0), Inches(6.65), Inches(11.333), Inches(0.6))
        tf = desc_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = f'"{description}"'
        p.font.name = FONT_BODY
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = COLOR_DARK_TEXT

        return slide

    # =========================================================================
    # SCREENSHOT SLIDES DEFINITIONS (33 Screenshot Slides starting Slide 23)
    # =========================================================================
    screenshot_slides_data = [
        # CUSTOMER APP
        ("Screenshot_2026-08-13-08-56-01-351_com.parksafe.parksafe.jpg",
         "Splash Screen",
         "App launch screen displaying ParkSafe logo and smart parking brand identity.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-56-06-677_com.parksafe.parksafe.jpg",
         "Onboarding - Find Nearby Parking",
         "First onboarding screen highlighting real-time nearby parking discovery.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-56-09-161_com.parksafe.parksafe.jpg",
         "Onboarding - Time Slot Booking",
         "Second onboarding screen showcasing flexible hourly and daily slot reservations.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-56-11-961_com.parksafe.parksafe.jpg",
         "Onboarding - Instant QR Access",
         "Third onboarding screen explaining paperless QR-based gate entry and exit.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-09-53-01-527_com.parksafe.parksafe.jpg",
         "Customer Registration",
         "Simple registration form for new users to create a secure ParkSafe account.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-56-15-381_com.parksafe.parksafe.jpg",
         "Customer Login & Quick Admin Selector",
         "Secure authentication screen with mobile password login and quick admin role switching.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-56-27-488_com.parksafe.parksafe.jpg",
         "City & Location Selection",
         "Allows customers to filter parking locations across major Indian cities.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-56-37-073_com.parksafe.parksafe.jpg",
         "Home Dashboard & AI Nearest Parking",
         "Main customer dashboard featuring smart parking recommendations based on proximity and real-time availability.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-56-45-351_com.parksafe.parksafe.jpg",
         "Parking Location Details & Google Maps Link",
         "Displays parking amenities, turn-by-turn Google Maps link, hourly rates, and vehicle type selection.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-56-48-322_com.parksafe.parksafe.jpg",
         "Date & Time Duration Selection",
         "Allows users to select check-in and check-out times, rental duration, and view estimated pricing.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-56-55-813_com.parksafe.parksafe.jpg",
         "Floor-Wise Slot Selection (Multi-Floor)",
         "Customers can view available parking slots organized floor by floor and select their preferred spot.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-56-58-040_com.parksafe.parksafe.jpg",
         "Multi-Floor Level Navigation Grid",
         "Interactive multi-floor layout showing real-time slot availability across all building levels.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-57-25-865_com.parksafe.parksafe.jpg",
         "Booking Confirmation Summary",
         "Displays assigned slot details, total duration, price breakdown, and vehicle plate entry.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-57-30-712_com.parksafe.parksafe.jpg",
         "Payment & NETC FASTag Auto-Scan",
         "Offers multiple payment options including NETC FASTag automatic RFID scan for contactless barrier entry.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-57-32-667_com.parksafe.parksafe.jpg",
         "Payment Methods & Pay at Counter",
         "Provides complete payment flexibility including digital gateways and physical counter payment.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-57-39-590_com.parksafe.parksafe.jpg",
         "Digital Booking Ticket",
         "Displays digital parking pass with assigned slot number, scheduled duration, and payment confirmation.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-57-41-570_com.parksafe.parksafe.jpg",
         "QR Ticket Actions & Google Maps Integration",
         "Enables customers to save ticket QR codes to gallery and launch turn-by-turn navigation.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-57-47-409_com.parksafe.parksafe.jpg",
         "My Bookings - Active Reservations",
         "Shows all active confirmed slot reservations with live countdown and location details.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-57-54-174_com.parksafe.parksafe.jpg",
         "My Bookings - Past & Cancelled Records",
         "Provides complete history of completed, expired, or cancelled slot bookings.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-57-58-046_com.parksafe.parksafe.jpg",
         "Customer Profile Management",
         "Allows customers to view personal information, manage refund requests, and access settings.",
         "CUSTOMER APP DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-58-03-031_com.parksafe.parksafe.jpg",
         "Customer Account Settings & Sign Out",
         "Enables users to update account details or securely log out.",
         "CUSTOMER APP DEMONSTRATION"),

        # ADMIN PANEL
        ("Screenshot_2026-08-13-08-26-01-730_com.parksafe.parksafe.jpg",
         "Admin Dashboard",
         "Provides real-time statistics on bookings, today's revenue, active checked-in vehicles, and available slots.",
         "ADMIN PANEL DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-26-15-622_com.parksafe.parksafe.jpg",
         "Admin QR Code Scanner",
         "Allows parking attendants to scan customer QR tickets or enter vehicle numbers manually for instant lookup.",
         "ADMIN PANEL DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-26-21-585_com.parksafe.parksafe.jpg",
         "Admin Today's Bookings",
         "Displays all bookings scheduled for the current day with instant status filters.",
         "ADMIN PANEL DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-26-31-819_com.parksafe.parksafe.jpg",
         "Admin All Bookings Management",
         "Allows administrators to track, view, and search complete historical and active booking records.",
         "ADMIN PANEL DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-26-34-471_com.parksafe.parksafe.jpg",
         "Admin Registered Vehicles Tracker",
         "Tracks frequency of vehicle visits and last recorded parking date for analytical tracking.",
         "ADMIN PANEL DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-26-40-566_com.parksafe.parksafe.jpg",
         "Admin Slot & Floor Management",
         "Enables staff to view active locations and configure slot availability across different floors.",
         "ADMIN PANEL DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-26-45-078_com.parksafe.parksafe.jpg",
         "Admin Add Location & Multi-Floor Configuration",
         "Allows administrators to add new parking locations, set hourly pricing, and configure multi-floor slot distribution.",
         "ADMIN PANEL DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-27-33-438_com.parksafe.parksafe.jpg",
         "Admin Customer Management",
         "Displays registered customer profiles, contact numbers, and total user records.",
         "ADMIN PANEL DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-27-42-125_com.parksafe.parksafe.jpg",
         "Admin Customer Booking History View",
         "Allows administrators to inspect complete booking history and slot details for any individual customer.",
         "ADMIN PANEL DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-27-49-170_com.parksafe.parksafe.jpg",
         "Admin Staff Account Overview",
         "Shows administrative staff credentials and location management permissions.",
         "ADMIN PANEL DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-28-02-580_com.parksafe.parksafe.jpg",
         "Admin Revenue Analytics",
         "Provides financial insights, earnings trends, and daily breakdown of parking revenue.",
         "ADMIN PANEL DEMONSTRATION"),

        ("Screenshot_2026-08-13-08-29-23-775_com.parksafe.parksafe.jpg",
         "Admin Refund Request Management",
         "Enables administrators to review customer refund claims and process approvals or rejections.",
         "ADMIN PANEL DEMONSTRATION"),
    ]

    for idx, (img_fn, title, desc, cat) in enumerate(screenshot_slides_data):
        print(f"Adding Slide {23 + idx}: {title}")
        add_screenshot_slide(prs, img_fn, title, desc, category=cat)

    # =========================================================================
    # FINAL SLIDE: REFERENCES & TECHNOLOGIES USED
    # =========================================================================
    print(f"Adding Final Slide {23 + len(screenshot_slides_data)}: References & Technologies Used")
    slide_ref = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide_ref)
    add_slide_header(slide_ref, "References & Technologies Used", category_text="PARKSAFE - PROJECT ARCHITECTURE & REFERENCES")

    tech_stack = [
        ("Flutter", "Cross-platform framework for Web, Android & iOS apps"),
        ("Dart", "Primary programming language for mobile application"),
        ("Supabase", "Cloud database, user authentication & realtime backend"),
        ("PostgreSQL", "Relational database engine with secure RLS policies"),
        ("Firebase FCM", "Push notification service for reminders & alerts"),
        ("Provider", "Reactive state management pattern for Flutter UI"),
        ("GoRouter", "Declarative routing and URL navigation system"),
        ("Geolocator", "GPS tracking to measure distance to parking"),
        ("Mobile Scanner", "High-speed camera QR code reader package"),
        ("Google Maps API", "Map location links and turn-by-turn navigation"),
        ("Python & FastAPI", "AI recommendation microservice REST API"),
        ("Scikit-learn", "Machine learning utilities for recommendation scoring"),
        ("Figma", "UI/UX screen design and logo branding assets"),
        ("Antigravity", "AI-assisted development workspace & code intelligence"),
        ("GitHub", "Version control and code repository management")
    ]

    for idx, (tech_name, tech_desc) in enumerate(tech_stack):
        r = idx // 3
        c = idx % 3
        left = Inches(0.8 + c * 3.9)
        top = Inches(1.5 + r * 1.08)

        shape = slide_ref.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, Inches(3.6), Inches(0.95))
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_CARD_BG
        shape.line.color.rgb = COLOR_BORDER
        shape.line.width = Pt(1.2)

        tbox = slide_ref.shapes.add_textbox(left + Inches(0.15), top + Inches(0.1), Inches(3.3), Inches(0.75))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = tech_name
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_PRIMARY_GREEN

        p1 = tf.add_paragraph()
        p1.text = tech_desc
        p1.font.name = FONT_BODY
        p1.font.size = Pt(11)
        p1.font.color.rgb = COLOR_DARK_TEXT
        p1.space_before = Pt(2)

    prs.save(output_pptx)
    print(f"Successfully generated final presentation with {len(prs.slides)} slides at: {output_pptx}")

if __name__ == '__main__':
    generate_presentation()
