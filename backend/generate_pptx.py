import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_parksafe_presentation():
    output_path = r'C:/Users/isida/Desktop/ParkSafe_College_Presentation.pptx'
    logo_path = r'C:/Users/isida/Desktop/Parksafe/assets/images/parksafe_logo_clean.png'
    watermark_path = r'C:/Users/isida/Desktop/Parksafe/assets/images/parksafe_watermark.png'

    prs = Presentation()
    prs.slide_width = Inches(13.333) # 16:9 Widescreen
    prs.slide_height = Inches(7.5)

    # Color Palette
    COLOR_WHITE = RGBColor(255, 255, 255)
    COLOR_PRIMARY_GREEN = RGBColor(15, 81, 50) # #0F5132
    COLOR_DARK_TEXT = RGBColor(20, 24, 27) # #14181B
    COLOR_MUTED_TEXT = RGBColor(91, 102, 96) # #5B6660
    COLOR_CARD_BG = RGBColor(248, 250, 249) # Subtle light tint card
    COLOR_BORDER = RGBColor(226, 232, 240) # Light border

    FONT_HEADING = 'Walbaum Display'
    FONT_BODY = 'Aptos Light'

    def set_white_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = COLOR_WHITE

    def add_watermark(slide):
        if os.path.exists(watermark_path):
            # Place subtle low-opacity logo in bottom right corner
            slide.shapes.add_picture(
                watermark_path,
                left=Inches(9.5),
                top=Inches(5.2),
                width=Inches(3.4)
            )

    def add_slide_header(slide, title_text, category_text="PARKSAFE PROJECT PRESENTATION"):
        # Header text frame
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.9))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        # Category / Kicker
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.name = FONT_BODY
        p0.font.size = Pt(10)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_MUTED_TEXT

        # Slide Title
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.name = FONT_HEADING
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_PRIMARY_GREEN

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide1)

    # Logo
    if os.path.exists(logo_path):
        slide1.shapes.add_picture(logo_path, Inches(0.8), Inches(0.8), width=Inches(3.8))

    # Title & Subtitle box
    tbox = slide1.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(2.2))
    tf1 = tbox.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = 0

    p_title = tf1.paragraphs[0]
    p_title.text = "ParkSafe"
    p_title.font.name = FONT_HEADING
    p_title.font.size = Pt(46)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_PRIMARY_GREEN

    p_sub = tf1.add_paragraph()
    p_sub.text = "Smart Vehicle Parking & Online Slot Booking Application"
    p_sub.font.name = FONT_BODY
    p_sub.font.size = Pt(22)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_DARK_TEXT

    p_tag = tf1.add_paragraph()
    p_tag.text = "FIND. BOOK. PARK SAFE."
    p_tag.font.name = FONT_BODY
    p_tag.font.size = Pt(14)
    p_tag.font.bold = True
    p_tag.font.color.rgb = COLOR_PRIMARY_GREEN

    # Team & Department Box (Clean Layout)
    team_box = slide1.shapes.add_textbox(Inches(0.8), Inches(4.8), Inches(5.5), Inches(2.0))
    tf_team = team_box.text_frame
    tf_team.word_wrap = True
    tf_team.margin_left = tf_team.margin_top = tf_team.margin_right = tf_team.margin_bottom = 0

    pt0 = tf_team.paragraphs[0]
    pt0.text = "TEAM MEMBERS:"
    pt0.font.name = FONT_BODY
    pt0.font.size = Pt(12)
    pt0.font.bold = True
    pt0.font.color.rgb = COLOR_PRIMARY_GREEN

    for name in ["Sayyed Mashhuda", "Sidat Ibrahim", "Patel Dhruvil"]:
        pt = tf_team.add_paragraph()
        pt.text = f"- {name}"
        pt.font.name = FONT_BODY
        pt.font.size = Pt(13)
        pt.font.color.rgb = COLOR_DARK_TEXT

    dept_box = slide1.shapes.add_textbox(Inches(6.8), Inches(4.8), Inches(5.5), Inches(2.0))
    tf_dept = dept_box.text_frame
    tf_dept.word_wrap = True
    tf_dept.margin_left = tf_dept.margin_top = tf_dept.margin_right = tf_dept.margin_bottom = 0

    pd0 = tf_dept.paragraphs[0]
    pd0.text = "DEPARTMENT:"
    pd0.font.name = FONT_BODY
    pd0.font.size = Pt(12)
    pd0.font.bold = True
    pd0.font.color.rgb = COLOR_PRIMARY_GREEN

    pd1 = tf_dept.add_paragraph()
    pd1.text = "B.Voc (Bachelor of Vocation)"
    pd1.font.name = FONT_BODY
    pd1.font.size = Pt(14)
    pd1.font.bold = True
    pd1.font.color.rgb = COLOR_DARK_TEXT

    pd2 = tf_dept.add_paragraph()
    pd2.text = "Software Development / Computer Applications"
    pd2.font.name = FONT_BODY
    pd2.font.size = Pt(12)
    pd2.font.color.rgb = COLOR_MUTED_TEXT


    # Helper to add content cards
    def add_card(slide, left, top, width, height, title, points, icon_symbol="📌"):
        # Card Background Shape
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = COLOR_CARD_BG
        shape.line.color.rgb = COLOR_BORDER
        shape.line.width = Pt(1)

        # Content Text Box
        tbox = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
        tf = tbox.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p0 = tf.paragraphs[0]
        p0.text = f"{icon_symbol} {title}"
        p0.font.name = FONT_HEADING
        p0.font.size = Pt(15)
        p0.font.bold = True
        p0.font.color.rgb = COLOR_PRIMARY_GREEN

        for pt in points:
            p = tf.add_paragraph()
            p.text = f"- {pt}"
            p.font.name = FONT_BODY
            p.font.size = Pt(12)
            p.font.color.rgb = COLOR_DARK_TEXT

    # =========================================================================
    # SLIDE 2: INTRODUCTION & PROJECT BACKGROUND
    # =========================================================================
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide2)
    add_watermark(slide2)
    add_slide_header(slide2, "1. Introduction & Project Background")

    add_card(slide2, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "Urbanization Crisis",
             ["Rapid growth in private vehicle ownership in urban centers.",
              "Limited availability of organized parking spaces.",
              "Severe traffic bottlenecks caused by drivers searching for open slots.",
              "Need for intelligent transportation & smart city infrastructure."], "🏙️")

    add_card(slide2, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "The Digital Shift",
             ["Transitioning from manual paper tokens to digital reservations.",
              "Eliminating uncertainty prior to starting a journey.",
              "Providing real-time visibility into multi-floor parking complexes.",
              "Seamless integration with GPS location services for turn-by-turn navigation."], "📱")

    add_card(slide2, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2),
             "The ParkSafe Vision",
             ["Unified cross-platform mobility ecosystem built on Flutter & Supabase.",
              "Intelligent AI recommendation engine for nearest available slots.",
              "Contactless check-in via secure QR pass verification.",
              "End-to-end automated slot tracking, overtime calculation & payment."], "🎯")

    # =========================================================================
    # SLIDE 3: PROBLEM STATEMENT
    # =========================================================================
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide3)
    add_watermark(slide3)
    add_slide_header(slide3, "2. Problem Statement: Urban Parking Crises")

    add_card(slide3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.4),
             "Time & Fuel Wastage",
             ["Studies show drivers spend an average of 15–20 minutes searching for parking.",
              "Accounts for up to 30% of unnecessary inner-city traffic congestion.",
              "Excess fuel consumption leading to increased carbon emissions."], "⛽")

    add_card(slide3, Inches(6.8), Inches(1.6), Inches(5.7), Inches(2.4),
             "Lack of Real-Time Information",
             ["Drivers arrive at garages blindly without knowing slot availability.",
              "No distinction between vehicle slot types (Car, Bike, EV Charging).",
              "Inefficient space utilization across multi-floor complexes."], "❌")

    add_card(slide3, Inches(0.8), Inches(4.3), Inches(5.6), Inches(2.5),
             "Manual Garage Operations",
             ["Paper ticket loss leads to administrative disputes and friction.",
              "Lack of automated overstay tracking and accurate fee collection.",
              "Human error in gate check-in/check-out auditing."], "📝")

    add_card(slide3, Inches(6.8), Inches(4.3), Inches(5.7), Inches(2.5),
             "Security & Verification Risks",
             ["Difficulty verifying ticket ownership without digital credentials.",
              "Potential for fraudulent ticket duplication.",
              "No automated notification alert system for upcoming expiry."], "🛡️")

    # =========================================================================
    # SLIDE 4: PROPOSED SOLUTION
    # =========================================================================
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide4)
    add_watermark(slide4)
    add_slide_header(slide4, "3. Proposed Solution: The ParkSafe Platform")

    add_card(slide4, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "Pre-Booking System",
             ["Reserve guaranteed parking slots ahead of arrival.",
              "Select exact date, check-in time, and check-out time.",
              "Filter slots specifically for Car, Bike, or EV charging.",
              "Single-page scrollable visual floor grid."], "📅")

    add_card(slide4, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "AI Recommendation",
             ["Haversine formula calculates distance from live GPS.",
              "Combines proximity + real-time slot availability ratio.",
              "Fast Python FastAPI backend scoring engine.",
              "Direct navigation via Google Maps links."], "⚡")

    add_card(slide4, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2),
             "Digital Gatekeeping",
             ["Contactless QR code ticket generation upon booking.",
              "Instant Admin QR scanner for entry check-in & exit check-out.",
              "Automatic overtime & penalty calculation upon checkout.",
              "Firebase FCM real-time push notifications."], "🎟️")

    # =========================================================================
    # SLIDE 5: OBJECTIVES & SCOPE
    # =========================================================================
    slide5 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide5)
    add_watermark(slide5)
    add_slide_header(slide5, "4. Key Objectives & Project Scope")

    add_card(slide5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Core Technical Objectives",
             ["Zero Double-Booking Guarantee: Enforce precise datetime overlap validation at database level.",
              "Multi-Floor Architecture: Support dynamic floor-level slot management (Floor 1, Floor 2, Basement).",
              "Real-Time Sync: Immediate cloud database synchronization between Customer app & Admin panel.",
              "High Performance: Sub-second slot search and AI recommendation calculations.",
              "Cross-Platform Deployment: Single Flutter codebase running seamlessly on Android, iOS & Web."], "🎯")

    add_card(slide5, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
             "Project Scope & Boundary",
             ["Customer App: Search, AI recommendation, vehicle slot selection, digital wallet payment, QR ticket.",
              "Admin Panel: Add locations, configure floor-wise slot counts, QR scanner, live dashboard, refund management.",
              "Backend Microservice: Python FastAPI service for geospatial scoring and automated push notification scheduler.",
              "Database Infrastructure: Supabase PostgreSQL with Row Level Security (RLS) policies."], "🔍")

    # =========================================================================
    # SLIDE 6: USER ROLES
    # =========================================================================
    slide6 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide6)
    add_watermark(slide6)
    add_slide_header(slide6, "5. User Roles & System Operations")

    add_card(slide6, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Customer (Vehicle Owner)",
             ["Account Registration & Secure Mobile/Password Login.",
              "AI-Powered Nearest & Available Parking Location Search.",
              "Date & Time Range Selection for Guaranteed Slot Booking.",
              "Floor & Slot Selection (Car, Bike, EV filtered).",
              "Digital QR Ticket Storage in 'My Bookings'.",
              "In-App & FCM Push Notifications for Reminders & Alerts.",
              "Cancellation & Refund Request Processing."], "👤")

    add_card(slide6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
             "Admin (Garage Operator)",
             ["Parking Location Creation & Google Maps Link Integration.",
              "Multi-Floor Slot Setup & Capacity Allocation per Vehicle Type.",
              "QR Code Camera Scanner for Entry Check-In & Exit Check-Out.",
              "Manual Vehicle License Plate Search for Rapid Lookup.",
              "Live Dashboard: Total Revenue, Occupancy Rate, Today's Bookings.",
              "Refund Approval / Rejection Audit Workflow."], "👨‍💼")

    # =========================================================================
    # SLIDE 7: MAIN FEATURES OVERVIEW
    # =========================================================================
    slide7 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide7)
    add_watermark(slide7)
    add_slide_header(slide7, "6. Main Core Features Overview")

    features = [
        ("Multi-Floor Management", "Dynamic floor setup (Floor 1, 2, Basement) with custom Car/Bike/EV slot limits.", "🏢"),
        ("AI Recommender", "Ranks nearest locations using Haversine distance + real-time slot availability score.", "🧠"),
        ("Datetime Overlap Check", "Database-level overlap query prevents any overlapping booking collisions.", "⏰"),
        ("QR Pass Verification", "Instant gate entry/exit verification with early check-in protection.", "📲"),
        ("Firebase FCM Alerts", "Automated 15-min check-in, 15-min check-out, and overtime warning alerts.", "🔔"),
        ("Overtime & Refunds", "Automatic penalty fee calculation for overstaying and transparent refund management.", "💰")
    ]

    for idx, (ftitle, fdesc, ficon) in enumerate(features):
        r = idx // 3
        c = idx % 3
        left = Inches(0.8 + c * 3.9)
        top = Inches(1.6 + r * 2.7)
        add_card(slide7, left, top, Inches(3.6), Inches(2.4), ftitle, [fdesc], ficon)

    # =========================================================================
    # SLIDE 8: CUSTOMER BOOKING PROCESS
    # =========================================================================
    slide8 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide8)
    add_watermark(slide8)
    add_slide_header(slide8, "7. Customer Booking Process & User Flow")

    steps = [
        ("Step 1: Location & Time", "Select parking location, booking date, check-in time, and check-out time.", "1️⃣"),
        ("Step 2: Vehicle Type", "Choose vehicle category: Car 🚗, Bike 🏍️, or EV Charging ⚡.", "2️⃣"),
        ("Step 3: Floor & Slot", "View available slots grouped by floor on a single scrollable grid.", "3️⃣"),
        ("Step 4: QR Ticket Pass", "Instant payment & QR ticket generated for digital gate check-in.", "4️⃣")
    ]

    for idx, (stitle, sdesc, sicon) in enumerate(steps):
        left = Inches(0.8 + idx * 2.95)
        add_card(slide8, left, Inches(1.8), Inches(2.75), Inches(4.8), stitle, [sdesc], sicon)

    # =========================================================================
    # SLIDE 9: MULTI-FLOOR PARKING
    # =========================================================================
    slide9 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide9)
    add_watermark(slide9)
    add_slide_header(slide9, "8. Multi-Floor Parking & Slot Management")

    add_card(slide9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Admin Multi-Floor Configuration",
             ["Multi-Floor Toggle: Enable multi-floor setup when creating a place.",
              "Dynamic Floor Forms: Admin enters number of floors (e.g. 5 floors).",
              "Independent Capacities: Set Car, Bike, and EV slots separately per floor.",
              "Strict Validation: Enforces Car + Bike + EV = Total Slots for each floor.",
              "Automated Slot Generation: Generates linked slots (e.g., C-01, E-01) with floor attributes."], "🏗️")

    add_card(slide9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
             "Customer Floor-Wise Layout",
             ["Single Scrollable Page: All floors displayed on one seamless screen.",
              "Floor Headings: Clear section headers (e.g. 'Floor 1', 'Floor 2', 'Basement 1').",
              "Live Availability Badges: Displays available count per floor.",
              "Vehicle Filter: Displays only requested vehicle type slots under each floor.",
              "Ticket Synchronization: Booking ticket & My Bookings clearly show 'Floor 2 → Slot C-24'."], "🏢")

    # =========================================================================
    # SLIDE 10: REAL-TIME AVAILABILITY
    # =========================================================================
    slide10 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide10)
    add_watermark(slide10)
    add_slide_header(slide10, "9. Real-Time Availability & Overlap Logic")

    add_card(slide10, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Database Overlap Formula",
             ["Queries active non-cancelled bookings for requested slot.",
              "Mathematical Overlap Condition: (NewStart < ExistEnd) AND (NewEnd > ExistStart).",
              "If overlap is detected, slot status is marked Unavailable.",
              "Cancelled bookings immediately release slot availability back to pool."], "🔄")

    add_card(slide10, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
             "Same-Day Non-Overlapping Bookings",
             ["Supports multiple bookings for the same slot on the same day.",
              "Example: Slot C-01 can be booked from 1:00 PM–2:00 PM and 2:00 PM–4:00 PM.",
              "Prevents partial time collisions (e.g., 1:30 PM–2:30 PM blocked).",
              "Guarantees 100% reservation accuracy with zero double bookings."], "✅")

    # =========================================================================
    # SLIDE 11: AI RECOMMENDATION ENGINE
    # =========================================================================
    slide11 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide11)
    add_watermark(slide11)
    add_slide_header(slide11, "10. AI Nearest & Available Recommendation")

    add_card(slide11, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "Geospatial Calculation",
             ["Uses Haversine distance formula based on GPS lat/lng.",
              "Calculates precise distance in kilometers from user location.",
              "Estimates ETA minutes assuming standard urban driving speeds."], "📍")

    add_card(slide11, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "Availability Scoring",
             ["Queries live slot occupancy for requested vehicle type.",
              "Calculates Availability Score = Free Slots / Total Slots.",
              "Filters out locations with zero available slots of requested vehicle type."], "📊")

    add_card(slide11, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2),
             "Weighted Hybrid Score",
             ["Formula: 0.60 * Distance Score + 0.40 * Availability Score.",
              "Ranks top recommendation results dynamically.",
              "Exposed via fast Python FastAPI REST endpoint."], "🤖")

    # =========================================================================
    # SLIDE 12: AI/ML & FRAUD DETECTION
    # =========================================================================
    slide12 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide12)
    add_watermark(slide12)
    add_slide_header(slide12, "11. AI/ML Analytics & Fraud Detection")

    add_card(slide12, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Predictive Peak Hour Analytics",
             ["Analyzes historical booking trends by location and hour.",
              "Provides percentage prediction insight for future time slots.",
              "Helps customers plan ahead for peak shopping or business hours."], "📈")

    add_card(slide12, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
             "Fraud Prevention & Security",
             ["Cryptographic QR Ticket Signature: Prevents QR code tampering.",
              "Duplicate Pass Guard: Prevents reuse of completed or cancelled passes.",
              "Vehicle License Match: Admin verifies vehicle number plate on check-in.",
              "Audit Trail: Logs check-in/out timestamps and admin IDs."], "🔒")

    # =========================================================================
    # SLIDE 13: QR CHECK-IN & CHECK-OUT
    # =========================================================================
    slide13 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide13)
    add_watermark(slide13)
    add_slide_header(slide13, "12. QR Check-In & Check-Out Verification")

    add_card(slide13, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Admin QR Scanner Workflow",
             ["Camera Scanner: Built-in high speed camera QR reader in Admin panel.",
              "Instant Lookup: Fetches customer profile, vehicle number & slot details.",
              "Manual Plate Search: Fallback search by vehicle number plate.",
              "Status Transition: Updates booking state to 'Checked In' or 'Completed'."], "📷")

    add_card(slide13, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
             "Early Check-In Rule Enforcement",
             ["No Early Check-In Rule: Prevents check-in prior to scheduled start time.",
              "Clear UI Feedback: Alerts admin if customer arrives early with countdown mins.",
              "Timestamp Precision: Stores exact actual check-in & check-out DateTime in UTC."], "⏱️")

    # =========================================================================
    # SLIDE 14: OVERTIME & PENALTY SYSTEM
    # =========================================================================
    slide14 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide14)
    add_watermark(slide14)
    add_slide_header(slide14, "13. Booking Status, Overtime & Penalty Fee")

    add_card(slide14, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Booking Lifecycle States",
             ["Pending: Booking created, awaiting arrival.",
              "Checked In: Vehicle scanned & parked inside garage.",
              "Completed: Vehicle checked out & total payment settled.",
              "Cancelled: Booking cancelled by user prior to check-in."], "🔄")

    add_card(slide14, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
             "Automated Overtime Fee Engine",
             ["Overtime Condition: Triggered when actual checkout > scheduled end time.",
              "Penalty Calculation: Extra Hours = ceil(overtime_minutes / 60).",
              "Fee Calculation: Overtime Fee = Extra Hours * Hourly Rate.",
              "Real-Time Sync: Overtime fee dynamically added to total invoice upon checkout."], "💳")

    # =========================================================================
    # SLIDE 15: FIREBASE PUSH NOTIFICATIONS
    # =========================================================================
    slide15 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide15)
    add_watermark(slide15)
    add_slide_header(slide15, "14. Firebase Push Notifications & Reminders")

    add_card(slide15, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "15-Min Check-In Reminder",
             ["Pushed 15 minutes before scheduled start time.",
              "Alerts customer to get ready and navigate to location.",
              "Delivered even when app is closed or in background."], "🔔")

    add_card(slide15, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "15-Min Check-Out Reminder",
             ["Pushed 15 minutes before scheduled end time.",
              "Reminds customer to return to vehicle or extend booking.",
              "Helps avoid unexpected overtime charges."], "⏰")

    add_card(slide15, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2),
             "Overtime & In-App Sync",
             ["Pushes warning alert if vehicle stays past end time.",
              "Saves notification record to Supabase database.",
              "Synchronized in real-time with customer Notification page."], "⚠️")

    # =========================================================================
    # SLIDE 16: PAYMENT & REFUND MANAGEMENT
    # =========================================================================
    slide16 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide16)
    add_watermark(slide16)
    add_slide_header(slide16, "15. Payment & Refund Management")

    add_card(slide16, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Payment Integration",
             ["Digital Payment Gateways: Supports Card, UPI, and Wallet payments.",
              "Instant Receipt: Generates transaction reference ID.",
              "Transparent Pricing: Displays clear breakdown of hourly rate * hours."], "💳")

    add_card(slide16, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
             "Refund Workflow",
             ["Customer Cancellation: Request refund before check-in.",
              "Admin Approval Dashboard: Admin reviews & processes refund.",
              "Refund Status Tracking: 'Requested', 'Approved', 'Refunded', 'Rejected'.",
              "Automated In-App Notification when refund status updates."], "💸")

    # =========================================================================
    # SLIDE 17: GOOGLE MAPS INTEGRATION
    # =========================================================================
    slide17 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide17)
    add_watermark(slide17)
    add_slide_header(slide17, "16. Google Maps & Location Navigation")

    add_card(slide17, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Google Maps Integration",
             ["Direct Navigation Link: Opens native Google Maps app with destination query.",
              "Admin Link Management: Admins can update Google Maps URL for any location.",
              "Distance & Route: Displays distance in km and estimated driving time."], "🗺️")

    add_card(slide17, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
             "Parking Amenities & Facilities",
             ["Custom Amenity Chips: 24/7 CCTV, EV Charging, Car Wash, Security Guard.",
              "Filter & Badges: Highlights premium features on location details page.",
              "Admin Facility Control: Easily toggle amenities when configuring place."], "🌟")

    # =========================================================================
    # SLIDE 18: TECHNOLOGY STACK
    # =========================================================================
    slide18 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide18)
    add_watermark(slide18)
    add_slide_header(slide18, "17. Technology Stack & Architecture")

    add_card(slide18, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "Frontend (Flutter)",
             ["Flutter SDK (Dart) for Web, Android & iOS.",
              "Provider State Management for reactive UI.",
              "Custom Dark & Light Theme system.",
              "Flutter Local Notifications & Mobile Scanner."], "💙")

    add_card(slide18, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2),
             "Cloud (Supabase)",
             ["Supabase PostgreSQL Relational Database.",
              "Row Level Security (RLS) policies.",
              "Supabase Auth & Storage integration.",
              "Real-time database listener channels."], "⚡")

    add_card(slide18, Inches(8.8), Inches(1.6), Inches(3.7), Inches(5.2),
             "Backend (Python)",
             ["FastAPI REST framework with Uvicorn.",
              "APScheduler for background notification worker.",
              "Firebase Admin SDK for FCM push notifications.",
              "NumPy & Haversine AI calculation engine."], "🐍")

    # =========================================================================
    # SLIDE 19: BENEFITS & FUTURE SCOPE
    # =========================================================================
    slide19 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide19)
    add_watermark(slide19)
    add_slide_header(slide19, "18. Project Benefits & Future Scope")

    add_card(slide19, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2),
             "Key Achievements & Benefits",
             ["Driver Convenience: Eliminates parking hunting time & guarantees slots.",
              "Operator Efficiency: Automates gate check-ins and revenue collection.",
              "Environmental Impact: Reduces fuel waste & urban carbon emissions.",
              "Eco-Friendly: 100% paperless digital QR tickets."], "🌱")

    add_card(slide19, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2),
             "Future Innovations",
             ["ANPR / ALPR Integration: Automatic License Plate Recognition at gate barriers.",
              "Dynamic Surge Pricing: AI price adjustment during peak demand hours.",
              "IoT Sensor Integration: Real-time ultrasonic slot detection hardware.",
              "EV Smart Charging Auto-Pay: Automated charging billing upon checkout."], "🚀")

    # =========================================================================
    # SLIDE 20: CONCLUSION & Q/A
    # =========================================================================
    slide20 = prs.slides.add_slide(prs.slide_layouts[6])
    set_white_background(slide20)
    add_watermark(slide20)

    # Big Centered Final Layout
    tbox20 = slide20.shapes.add_textbox(Inches(1.5), Inches(1.5), Inches(10.3), Inches(4.5))
    tf20 = tbox20.text_frame
    tf20.word_wrap = True

    p_c0 = tf20.paragraphs[0]
    p_c0.alignment = PP_ALIGN.CENTER
    p_c0.text = "PARKSAFE PROJECT PRESENTATION"
    p_c0.font.name = FONT_BODY
    p_c0.font.size = Pt(14)
    p_c0.font.bold = True
    p_c0.font.color.rgb = COLOR_MUTED_TEXT

    p_c1 = tf20.add_paragraph()
    p_c1.alignment = PP_ALIGN.CENTER
    p_c1.text = "Thank You!"
    p_c1.font.name = FONT_HEADING
    p_c1.font.size = Pt(44)
    p_c1.font.bold = True
    p_c1.font.color.rgb = COLOR_PRIMARY_GREEN

    p_c2 = tf20.add_paragraph()
    p_c2.alignment = PP_ALIGN.CENTER
    p_c2.text = "Smart Vehicle Parking & Online Slot Booking Application"
    p_c2.font.name = FONT_BODY
    p_c2.font.size = Pt(18)
    p_c2.font.color.rgb = COLOR_DARK_TEXT

    p_c3 = tf20.add_paragraph()
    p_c3.alignment = PP_ALIGN.CENTER
    p_c3.text = "\nQuestions & Evaluation"
    p_c3.font.name = FONT_HEADING
    p_c3.font.size = Pt(22)
    p_c3.font.bold = True
    p_c3.font.color.rgb = COLOR_PRIMARY_GREEN

    p_c4 = tf20.add_paragraph()
    p_c4.alignment = PP_ALIGN.CENTER
    p_c4.text = "Department of B.Voc (Bachelor of Vocation) | Software Development"
    p_c4.font.name = FONT_BODY
    p_c4.font.size = Pt(13)
    p_c4.font.color.rgb = COLOR_MUTED_TEXT

    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_parksafe_presentation()
